"""Finance workbook entry point shared by the desktop workspace and CLI."""
from dataclasses import dataclass
from datetime import date
import os
from pathlib import Path
import tempfile
from openpyxl import load_workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900
from progress_studio.infrastructure.excel.finance_input_workbook import (
    INPUT, TABLES, capacity, check_schema, create_inputs, read_finance_inputs,
)
from progress_studio.infrastructure.excel.financial_forecast_workbook import build_finance_view
from progress_studio.infrastructure.excel.finance_package import merge_finance_parts
from progress_studio.infrastructure.excel.xlsx_package_validator import validate_xlsx_tables


@dataclass(frozen=True)
class FinanceWorkbookAnalysis:
    workbook: Path
    existing_finance: bool
    capacity_rows: int
    opening_date: date | None = None
    actuals_through: date | None = None
    currency: str | None = None


class FinancialForecastWorkbookService:
    def analyze(self, source_workbook) -> FinanceWorkbookAnalysis:
        """Read desktop readiness without changing the workbook or inferring cash.

        Main identifies the application workbook; its progress/BOQ readiness does
        not determine finance completeness. Existing finance uses FF-2 validation.
        """
        source = Path(source_workbook).expanduser().resolve()
        if source.suffix.lower() != '.xlsx':
            raise ValueError('Finance V1 requires .xlsx workbooks')
        if not source.is_file():
            raise FileNotFoundError(f'Workbook was not found: {source}')
        wb = load_workbook(source, keep_links=True)
        try:
            if 'main' not in wb.sheetnames:
                raise ValueError("Select a Progress Studio workbook containing 'main'.")
            if wb.epoch != CALENDAR_WINDOWS_1900:
                raise ValueError('Finance V1 requires the Excel 1900 date system')
            if not check_schema(wb):
                return FinanceWorkbookAnalysis(source, False, 100)
            inputs = read_finance_inputs(wb)
            return FinanceWorkbookAnalysis(
                source, True, max(capacity(wb[INPUT], name) for name in TABLES),
                inputs.opening_date, inputs.actuals_through, inputs.currency,
            )
        finally:
            wb.close()

    def generate(self, source_workbook, output_workbook, *, opening_date=None,
                 actuals_through=None, currency=None, capacity_rows=None):
        source=Path(source_workbook).expanduser().resolve()
        output=Path(output_workbook).expanduser().resolve()
        if source.suffix.lower()!='.xlsx' or output.suffix.lower()!='.xlsx':
            raise ValueError('Finance V1 requires .xlsx workbooks')
        if source==output or output.exists():raise ValueError('Finance output must be a new workbook path')
        if not source.is_file():raise FileNotFoundError(source)
        output.parent.mkdir(parents=True,exist_ok=True)
        fd,name=tempfile.mkstemp(prefix='.finance.',suffix='.xlsx',dir=output.parent)
        os.close(fd);temp=Path(name)
        try:
            wb=load_workbook(source,keep_links=True)
            try:
                if wb.epoch != CALENDAR_WINDOWS_1900:
                    raise ValueError("Finance V1 requires the Excel 1900 date system")
                create_inputs(wb,opening_date=opening_date,actuals_through=actuals_through,
                              currency=currency,capacity_rows=capacity_rows if capacity_rows is not None else 1 if "Finance Input" in wb.sheetnames else 100)
                read_finance_inputs(wb)  # FF-1 reference validation; never save invalid edits.
                build_finance_view(wb)
                wb.save(temp)
            finally:wb.close()
            merge_finance_parts(source,temp)
            validate_xlsx_tables(temp)
            # Reopen the published package, not only the intermediate serializer.
            check=load_workbook(temp,keep_links=True)
            try:read_finance_inputs(check)
            finally:check.close()
            os.replace(temp,output)
        finally:temp.unlink(missing_ok=True)
        return output
