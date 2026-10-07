"""Monthly presentation on Weekly geometry with the full Weekly overlay."""
from __future__ import annotations

from copy import copy, deepcopy
from datetime import date, datetime

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Protection, Side
from openpyxl.utils import get_column_letter

from progress_studio.infrastructure.excel.traditional_overlay_workbook import (
    add_compact_weekly_overlay,
)
from progress_studio.infrastructure.excel.weekly_compact_monthly_workbook import project_monthly_regions

SHEET_NAME = "Weekly Compact"
PERIOD_WIDTH = 2.5


def build_weekly_compact(workbook, dataset, *, snapshot=False, value_source=None):
    """Project main at identical coordinates after traditional overlays exist.

    Information cells follow main. Monthly body/footer consume main_monthly;
    Snapshot uses finalized values. Charts consume the existing Weekly hybrid/
    Live source, with Compact-only checkpoint helpers. No progress calculation.
    """
    source = workbook["main"]
    header = dataset.header_row
    columns = [c for c in range(1, source.max_column + 1)
               if isinstance(source.cell(header, c).value, (date, datetime))]
    if not columns or not dataset.periods:
        return False
    if snapshot and value_source is None:
        raise ValueError("Weekly Compact Snapshot requires cached main values.")
    if "PS_WEEKLY_OVERLAY_CUTOFF" not in workbook.defined_names:
        raise ValueError("Finalize the Weekly cutoff/overlays before Weekly Compact.")
    if "main_monthly" not in workbook.sheetnames:
        raise ValueError("Finalize main_monthly before Weekly Compact.")
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]
    target = workbook.copy_worksheet(source)
    target.title = SHEET_NAME
    workbook._sheets.remove(target)
    workbook._sheets.insert(workbook.worksheets.index(source) + 1, target)
    target.conditional_formatting = deepcopy(source.conditional_formatting)
    target.auto_filter = deepcopy(source.auto_filter)
    target.sheet_properties = deepcopy(source.sheet_properties)
    target.sheet_view.showGridLines = False
    target.freeze_panes = source.freeze_panes
    target.cell(1, 1, "Activity Data — Weekly Compact")

    for row in target.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.protection = Protection(locked=True, hidden=cell.protection.hidden)
            if cell.row > header:
                src = source.cell(cell.row, cell.column)
                if snapshot:
                    cell.value = value_source.cell(cell.row, cell.column).value
                else:
                    ref = f"'main'!{src.coordinate}"
                    cell.value = f'=IF({ref}="","",{ref})'

    # Normalize the copied control before the monthly projection relocates it.
    for row in range(header + 1, source.max_row + 1):
        if source.cell(row, 12).value == "Cutoff Date":
            target.cell(row, 13, "=PS_WEEKLY_OVERLAY_CUTOFF")
            target.cell(row, 13).number_format = "dd/mm/yyyy"
            target.cell(row, 13).comment = None

    for i, col in enumerate(columns):
        letter = get_column_letter(col)
        target.column_dimensions[letter].width = PERIOD_WIDTH
        # Retained periods must not collapse category geometry.
        target.column_dimensions[letter].hidden = False
        current = source.cell(header, col).value
        previous = source.cell(header, columns[i-1]).value if i else None
        following = source.cell(header, columns[i+1]).value if i+1 < len(columns) else None
        new_month = previous is None or (current.year, current.month) != (previous.year, previous.month)
        end_month = following is None or (current.year, current.month) != (following.year, following.month)
        new_year = previous is None or current.year != previous.year
        for row in range(1, target.max_row + 1):
            cell = target.cell(row, col)
            border = copy(cell.border)
            border.left = Side(style="medium" if new_year else "thin", color="9CAFC0") if new_month else Side()
            border.right = Side(style="thin", color="9CAFC0") if end_month else Side()
            cell.border = border
            if row >= header - 1:
                cell.number_format = ";;;"
        if new_month and not isinstance(target.cell(2, col), MergedCell):
            target.cell(2, col, current.strftime("%b"))

    project_monthly_regions(workbook, dataset, target, snapshot=snapshot)

    # copy_worksheet does not copy validations or drawings. Do not add editable
    # controls or recreate the source's global names. New chart, shared sources.
    add_compact_weekly_overlay(workbook, dataset, target)
    return True
