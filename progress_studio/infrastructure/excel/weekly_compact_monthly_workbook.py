"""Monthly presentation of finalized main_monthly outputs on Weekly geometry."""
from copy import copy, deepcopy
from datetime import date, datetime

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Protection
from openpyxl.utils import get_column_letter

from progress_studio.infrastructure.excel.live_monthly_workbook import _display_month_buckets


def project_monthly_regions(workbook, dataset, target, *, snapshot=False):
    """Reuse monthly outputs, including their blank/zero and source-mode semantics."""
    source = workbook["main"]
    monthly = workbook["main_monthly"]
    buckets = _display_month_buckets(source)
    month_columns = {
        (cell.value.year, cell.value.month): cell.column
        for cell in monthly[dataset.header_row]
        if isinstance(cell.value, (date, datetime))
    }
    body = [r for row in dataset.rows
            if row.row_type.strip().lower() in {"activity", "wbs", "project summary"}
            and row.pa.strip().upper() == "P"
            for r in (row.row_number, row.row_number + 1)]
    footer = [row.row_number for row in dataset.rows
              if row.row_type.strip().lower() == "s-curve"
              and row.pa.strip().upper() in {"P", "AP", "A", "AA"}]
    rows = sorted(set(body + footer))
    # The existing monthly renderer owns band colors/conditions. Re-anchor its
    # nonblank test to the merged month's top-left while keeping row metadata.
    first_col = buckets[0][1][0]
    first_row = min(body) if body else dataset.header_row + 1
    old_anchor = f"{get_column_letter(first_col)}{first_row}"
    monthly_rules = [rule for rules in monthly.conditional_formatting._cf_rules.values()
                     for rule in rules if rule.type == "expression"
                     and any(old_anchor + '<>""' in f for f in rule.formula)]
    preserved = ConditionalFormattingList()
    for area, rules in target.conditional_formatting._cf_rules.items():
        if all(r.max_col < first_col for r in area.sqref.ranges):
            for rule in rules:
                preserved.add(str(area.sqref), deepcopy(rule))
    for key, columns in buckets:
        monthly_col = month_columns[key]
        first, last = columns[0], columns[-1]
        for row in rows:
            original_left = copy(target.cell(row, first).border)
            original_right = copy(target.cell(row, last).border)
            origin = monthly.cell(row, monthly_col)
            if snapshot and origin.data_type == "f":
                raise ValueError("Weekly Compact Snapshot requires finalized monthly values.")
            for col in columns:
                target.cell(row, col).value = None
                target.cell(row, col).fill = copy(origin.fill)
            if first != last:
                target.merge_cells(start_row=row, end_row=row, start_column=first, end_column=last)
            anchor = target.cell(row, first)
            ref = f"'main_monthly'!{origin.coordinate}"
            anchor.value = origin.value if snapshot else f'=IF({ref}="","",{ref})'
            anchor.font = copy(origin.font)
            anchor.fill = copy(origin.fill)
            anchor.number_format = origin.number_format
            anchor.alignment = Alignment(horizontal="center", vertical="center", shrinkToFit=False)
            anchor.protection = Protection(locked=True)
            border = copy(original_left)
            border.right = copy(original_right.right)
            anchor.border = border
            for col in columns:
                edge = copy(target.cell(row, col).border)
                edge.top, edge.bottom = copy(original_left.top), copy(original_left.bottom)
                target.cell(row, col).border = edge
        if body:
            for rule in monthly_rules:
                rule = deepcopy(rule)
                rule.formula = [f.replace(old_anchor + '<>""',
                                f'${get_column_letter(first)}{first_row}<>""') for f in rule.formula]
                preserved.add(f"{get_column_letter(first)}{min(body)}:{get_column_letter(last)}{max(body)}", rule)
    for priority, rule in enumerate((r for rules in preserved._cf_rules.values() for r in rules), 1):
        rule.priority = priority
    target.conditional_formatting = preserved

    # Preserve the only input on main. Compact displays the same name after the
    # actual generated footer, rather than leaving a copied mid-section control.
    controls = [row for row in range(dataset.header_row + 1, source.max_row + 1)
                if source.cell(row, 12).value == "Cutoff Date"]
    for row in controls:
        for col in (12, 13):
            if col >= first_col:
                continue  # Legacy narrow schemas already projected these cells.
            target.cell(row, col).value = None
            target.cell(row, col)._style = copy(source.cell(row, 2)._style)
            target.cell(row, col).protection = Protection(locked=True)
    if footer and controls:
        row = max(footer + body) + 1  # Legacy layouts may place footer above body.
        for dest, src in ((2, 12), (3, 13)):
            target.cell(row, dest)._style = copy(source.cell(controls[0], src)._style)
            target.cell(row, dest).protection = Protection(locked=True)
        target.cell(row, 2, "Cutoff Date")
        target.cell(row, 3, "=PS_WEEKLY_OVERLAY_CUTOFF")
        target.cell(row, 3).number_format = "dd/mm/yyyy"
        target.row_dimensions[row].height = target.row_dimensions[max(footer)].height or 20
