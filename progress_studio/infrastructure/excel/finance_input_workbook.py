"""Persistent FF-2 input adapter; Excel addresses do not escape this layer."""
from copy import copy, deepcopy
from datetime import date, datetime
from decimal import Decimal
import re
from openpyxl.styles import Font, Protection, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as col
from openpyxl.utils.cell import range_boundaries
from progress_studio.domain.financial_forecast import CashCategory, CashDirection, CashEvent, FinanceInputs, ForecastItem, ForecastOverride
from progress_studio.services.financial_forecast_deriver import FinancialForecastDeriver, FinanceValidationError, certificate_items
from progress_studio.infrastructure.excel.dashboard_workbook import _FONT
from progress_studio.infrastructure.excel.progress_workbook import replace_defined_name

INPUT, DATA, VIEW = 'Finance Input', 'Finance_Data', 'Cash Flow'
SCHEMA = 'Progress Studio Finance v1'
FIRST = 25
SETTINGS = dict(zip(('Currency','Opening_Date','Opening_Cash','Actuals_Through','Credit_Days','Retention_Rate','Actuals_Complete','Receivables_Complete','Cash_Out_Complete','Unlinked_Reconciled','Horizon_End'),range(4,15)))
TABLES = {
 'FF_Cash': (1, ('Event ID','Cash date','Amount','Direction','Category','Item ID','Correction reason','Notes')),
 'FF_Items': (10, ('Item ID','Amount','Direction','Expected / certification date','Category','Balance date','Retention','Kind','Retention release','Notes')),
 'FF_Overrides': (21, ('Item ID','Remaining amount','Cash date','Clear date','Reason','Notes')),
}
DATE_FORMAT = 'dd-mmm-yyyy'
MONEY_FORMAT = '#,##0.00;[Red](#,##0.00);"–"'


def excel_date(value, optional=False):
    if optional and value in (None, ''): return None
    if isinstance(value, datetime) and value.time() == datetime.min.time(): value=value.date()
    if type(value) is not date or value < date(1900,3,1):
        raise FinanceValidationError('Finance dates must be calendar dates from 1-Mar-1900')
    return value


def _id(value):
    if not isinstance(value,str) or not re.fullmatch(r'[A-Za-z0-9_.:\-]+',value):
        raise FinanceValidationError('Finance IDs use letters, digits, underscore, dot, colon or hyphen')
    return value


def _amount(value, optional=False):
    if optional and value in (None,''): return None
    try:
        result=Decimal(str(value))
        if isinstance(value,bool) or not result.is_finite() or abs(result)>Decimal('1e12'): raise ValueError
        return result
    except Exception as exc:
        raise FinanceValidationError('Excel finance amounts must be numeric within +/-1e12') from exc


def _boolean(value):
    # Excel validation lists can supply text booleans; do not use bool('FALSE').
    if value in ('TRUE','FALSE'): return value=='TRUE'
    if type(value) is not bool: raise FinanceValidationError('Confirmations must be TRUE or FALSE')
    return value


def capacity(ws,name): return range_boundaries(ws.tables[name].ref)[3]-FIRST+1


def check_schema(wb):
    if INPUT not in wb.sheetnames:
        if any(s in wb.sheetnames for s in (DATA,VIEW)): raise FinanceValidationError('Finance sheet ownership collision')
        return False
    ws=wb[INPUT]
    if ws['A1'].value!=SCHEMA: raise FinanceValidationError('Unrecognized Finance Input schema; no automatic probe migration')
    for name,(start,headers) in TABLES.items():
        if name not in ws.tables or capacity(ws,name)<1: raise FinanceValidationError('Finance input table missing or damaged')
        left,top,right,bottom=range_boundaries(ws.tables[name].ref)
        if (left,top,right)!=(start,FIRST-1,start+len(headers)-1): raise FinanceValidationError('Finance table layout changed')
        if tuple(ws.cell(FIRST-1,c).value for c in range(left,right+1))!=headers: raise FinanceValidationError('Finance table headers changed')
        for row in ws.iter_rows(min_row=bottom+1,min_col=left,max_col=right):
            if any(c.value is not None for c in row): raise FinanceValidationError('Finance data outside prepared table capacity')
    for name in (DATA,VIEW):
        if name in wb.sheetnames and wb[name]['A1'].value!=SCHEMA: raise FinanceValidationError('Finance output ownership collision')
    return True


def input_rows(ws,name):
    start,headers=TABLES[name]
    for row in ws.iter_rows(min_row=FIRST,max_row=FIRST+capacity(ws,name)-1,min_col=start,max_col=start+len(headers)-1):
        values=tuple(c.value for c in row)
        if any(v is not None for v in values):
            if any(c.data_type=='f' for c in row): raise FinanceValidationError('Finance inputs must be values, not formulas')
            yield values


def read_finance_inputs(wb):
    if not check_schema(wb): raise FinanceValidationError('Finance Input is not present')
    ws=wb[INPUT]; s={key:ws.cell(row,2).value for key,row in SETTINGS.items()}
    if any(ws.cell(row,2).data_type=='f' for row in SETTINGS.values()): raise FinanceValidationError('Finance settings must be values')
    if type(s['Credit_Days']) is not int or not 0<=s['Credit_Days']<=36500: raise FinanceValidationError('Credit days must be integer 0..36500')
    rate=_amount(s['Retention_Rate'])
    if not 0<=rate<=1: raise FinanceValidationError('Retention must be 0%..100%')
    events=tuple(CashEvent(_id(i),excel_date(d),_amount(a),CashDirection(dr),CashCategory(cat),_id(link) if link else None,reason or '') for i,d,a,dr,cat,link,reason,note in input_rows(ws,'FF_Cash'))
    items=[]
    for i,a,dr,d,cat,bal,ret,kind,release,note in input_rows(ws,'FF_Items'):
        _id(i); amount=_amount(a)
        if kind=='certificate':
            if dr!='in' or cat!='operating' or bal is not None or (ret is not None and _boolean(ret)):
                raise FinanceValidationError('Certificates are original gross operating receivables')
            items.extend(certificate_items(i,amount,excel_date(d),credit_days=s['Credit_Days'],retention_rate=rate,retention_release_date=excel_date(release,True)))
        elif kind=='manual':
            if release is not None: raise FinanceValidationError('Manual items use Expected date, not Retention release')
            items.append(ForecastItem(i,amount,CashDirection(dr),excel_date(d,True),CashCategory(cat),excel_date(bal,True),False if ret is None else _boolean(ret)))
        else: raise FinanceValidationError('Item Kind must be manual or certificate')
    overrides=tuple(ForecastOverride(_id(i),reason or '',_amount(a,True),excel_date(d,True),False if clear is None else _boolean(clear)) for i,a,d,clear,reason,note in input_rows(ws,'FF_Overrides'))
    for records,key in ((events,'event_id'),(items,'item_id'),(overrides,'item_id')):
        ids=[getattr(r,key).casefold() for r in records]
        if len(ids)!=len(set(ids)): raise FinanceValidationError('Finance IDs must be unique ignoring case')
    from dataclasses import replace
    ids={i.item_id.casefold():i.item_id for i in items}
    events=tuple(replace(e,item_id=ids.get(e.item_id.casefold(),e.item_id)) if e.item_id else e for e in events)
    overrides=tuple(replace(o,item_id=ids.get(o.item_id.casefold(),o.item_id)) for o in overrides)
    if not isinstance(s['Currency'],str):raise FinanceValidationError('Currency must be text')
    result=FinanceInputs(s['Currency'],excel_date(s['Opening_Date']),_amount(s['Opening_Cash']),excel_date(s['Actuals_Through']),events,tuple(items),overrides,*(_boolean(s[k]) for k in ('Actuals_Complete','Receivables_Complete','Cash_Out_Complete','Unlinked_Reconciled')),excel_date(s['Horizon_End'],True))
    FinancialForecastDeriver().derive(result)
    return result


def bind_names(wb):
    for key,row in SETTINGS.items():
        name='FF_'+key;target=f"'{INPUT}'!$B${row}";old=wb.defined_names.get(name)
        if old is not None and old.attr_text!=target: raise FinanceValidationError('Finance name collision: '+name)
        replace_defined_name(wb,name,target)


def style_sheet(ws):
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='B25' if ws.title==INPUT else 'B19'
    ws.sheet_state='hidden' if ws.title==DATA else 'visible'
    ws['A1']=SCHEMA;ws.row_dimensions[1].hidden=True
    for row in ws:
        for c in row:
            c.font=Font(name=_FONT,size=10,color='172B4D');c.alignment=Alignment(vertical='center')
    ws.protection.sheet=True;ws.protection.autoFilter=False
    ws.protection.selectLockedCells=False;ws.protection.selectUnlockedCells=False


def create_inputs(wb,*,opening_date,actuals_through,currency,capacity_rows=100):
    if type(capacity_rows) is not int or not 1<=capacity_rows<=2000: raise FinanceValidationError('Input capacity must be 1..2000')
    if check_schema(wb):
        ws=wb[INPUT]
        for name,(start,headers) in TABLES.items():
            old=capacity(ws,name)
            if capacity_rows<=old:continue
            for r in range(FIRST+old,FIRST+capacity_rows):
                for c in range(start,start+len(headers)):
                    ws.cell(r,c)._style=copy(ws.cell(FIRST+old-1,c)._style)
                    ws.cell(r,c).protection=Protection(locked=False)
            ws.tables[name].ref=f'{col(start)}{FIRST-1}:{col(start+len(headers)-1)}{FIRST+capacity_rows-1}'
            if ws.tables[name].autoFilter:ws.tables[name].autoFilter.ref=ws.tables[name].ref
            for dv in ws.data_validations.dataValidation:
                for region in list(dv.ranges):
                    if region.min_row==FIRST and start<=region.min_col<=start+len(headers)-1:
                        dv.add(f'{col(region.min_col)}{FIRST+old}:{col(region.max_col)}{FIRST+capacity_rows-1}')
        bind_names(wb);return ws
    excel_date(opening_date);excel_date(actuals_through)
    ws=wb.create_sheet(INPUT)
    ws['A2']='Finance inputs'
    defaults=(currency,opening_date,0,actuals_through,30,0.05,False,False,False,False,None)
    for (key,row),value in zip(SETTINGS.items(),defaults): ws.cell(row,1,key.replace('_',' '));ws.cell(row,2,value)
    notes=['Amounts: one project currency, consistent net cash / tax basis.',
           'BOQ Amount = Contract Value. Cash Out is an AC proxy.',
           'Opening cash is immediately before events on Opening Date.',
           'Finance actuals completeness is independent of Progress and EV dates.',
           'Credit days and retention apply only to gross certificates.',
           'Certificate links: ID:net or ID:retention. Release date is explicit.',
           'Blank override keeps base. Zero cancels. A reason is required.',
           'Snapshot amount is remaining at END of Balance date.',
           'Use prepared rows. Structural expansion requires Finance refresh.',
           'F9 / Save updates finance. Reconcile before confirming completeness.',
           'Terminal cash is not profit. Financing is separate from operating cash.',
           'IDs: letters, digits, underscore, dot, colon, hyphen; unique ignoring case.',
           'Dates from 1-Mar-1900; amounts within +/-1e12 for Excel precision.']
    for r,note in enumerate(notes,4): ws.cell(r,4,note)
    for name,(start,headers) in TABLES.items():
        for c,h in enumerate(headers,start): ws.cell(FIRST-1,c,h)
        table=Table(displayName=name,ref=f'{col(start)}{FIRST-1}:{col(start+len(headers)-1)}{FIRST+capacity_rows-1}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
        for r in range(FIRST,FIRST+capacity_rows):
            for c in range(start,start+len(headers)): ws.cell(r,c).protection=Protection(locked=False)
    for c in range(1,27): ws.column_dimensions[col(c)].width=21
    ws.column_dimensions['A'].width=28;ws.column_dimensions['B'].width=22
    ws.row_dimensions[FIRST-1].height=36;style_sheet(ws)
    for row in SETTINGS.values(): ws.cell(row,2).protection=Protection(locked=False)
    for row in (5,7,14): ws.cell(row,2).number_format=DATE_FORMAT
    ws['B6'].number_format=MONEY_FORMAT;ws['B9'].number_format='0.0%'
    for r in range(FIRST,FIRST+capacity_rows):
        for c in (2,13,15,18,23): ws.cell(r,c).number_format=DATE_FORMAT
        for c in (3,11,22): ws.cell(r,c).number_format=MONEY_FORMAT
    for coords,values in [('D','in,out'),('E','operating,financing'),('L','in,out'),('N','operating,financing'),('P','TRUE,FALSE'),('Q','manual,certificate'),('X','TRUE,FALSE')]:
        dv=DataValidation(type='list',formula1='"'+values+'"',allow_blank=True);dv.showErrorMessage=True;dv.errorStyle='stop';ws.add_data_validation(dv);dv.add(f'{coords}{FIRST}:{coords}{FIRST+capacity_rows-1}')
    for r in (10,11,12,13):
        dv=DataValidation(type='list',formula1='"TRUE,FALSE"');ws.add_data_validation(dv);dv.add(ws.cell(r,2))
    for row in ws:
        for c in row:
            if not c.protection.locked: c.font=Font(name=_FONT,size=10,color='0000FF')
    bind_names(wb);return ws


def preserve_finance_inputs(source,target):
    """Preserve raw edits, including unfinished input, with Mapping source precedence."""
    if not check_schema(source): return
    check_schema(target);src=source[INPUT]
    if INPUT in target.sheetnames:
        index=target.sheetnames.index(INPUT);del target[INPUT];dst=target.create_sheet(INPUT,index)
    else: dst=target.create_sheet(INPUT)
    for row in src:
        for c in row:
            out=dst.cell(c.row,c.column,c.value)
            out.font=copy(c.font);out.fill=copy(c.fill);out.border=copy(c.border)
            out.alignment=copy(c.alignment);out.protection=copy(c.protection)
            out.number_format=c.number_format
            if c.comment:out.comment=copy(c.comment)
            if c.hyperlink:out._hyperlink=copy(c.hyperlink)
    for k,v in src.column_dimensions.items():dst.column_dimensions[k]=copy(v)
    for k,v in src.row_dimensions.items():dst.row_dimensions[k]=copy(v)
    dst.sheet_view.showGridLines=src.sheet_view.showGridLines;dst.freeze_panes=src.freeze_panes
    dst.protection=copy(src.protection);dst.data_validations=copy(src.data_validations)
    for table in src.tables.values():dst.add_table(deepcopy(table))
    bind_names(target)
    from progress_studio.infrastructure.excel.financial_forecast_workbook import build_finance_view
    build_finance_view(target)
