"""Live FF-2 formulas. Dated cash owns funding; months only aggregate streams.

No Python result is cached into this view. Formula ranges are bounded by the
prepared input tables. No volatile functions or dynamic-array Excel dependency.
"""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as col
from progress_studio.infrastructure.excel.finance_input_workbook import (
    INPUT, DATA, VIEW, SCHEMA, FIRST, SETTINGS, TABLES, capacity, check_schema,
    bind_names, style_sheet, DATE_FORMAT, MONEY_FORMAT,
)
from progress_studio.services.financial_forecast_deriver import FinanceValidationError
from progress_studio.infrastructure.excel.dashboard_workbook import _FONT, _COLORS


def _ref(c,r): return f"'{INPUT}'!{c}{r}"
def _rng(c,n): return f"'{INPUT}'!${c}${FIRST}:${c}${FIRST+n-1}"
def _dr(c,n): return f'${c}${FIRST}:${c}${FIRST+n-1}'
def _true(x): return f'OR({x}=TRUE,{x}="TRUE")'
def _bool(x,blank=False):
    s=f'OR(ISLOGICAL({x}),{x}="TRUE",{x}="FALSE")'
    return f'OR({x}="",{s})' if blank else s
def _date(x,optional=False):
    s=f'AND(ISNUMBER({x}),{x}>=61,{x}<=2958465,MOD({x},1)=0)'
    return f'OR({x}="",{s})' if optional else s

def _money(x,signed=False,optional=False):
    s=f'AND(ISNUMBER({x}),ABS({x})<=1E+12'+('' if signed else f',{x}>=0')+')'
    return f'OR({x}="",{s})' if optional else s

def _id(x):
    # SUMIFS/MATCH wildcard syntax must not reinterpret a business identity.
    return f'AND(ISTEXT({x}),{x}<>"",{x}=TRIM({x}),ISERROR(FIND("*",{x})),ISERROR(FIND("?",{x})),ISERROR(FIND("~",{x})))'

def _set(ws,c,r,expr): ws.cell(r,c,'='+expr)

def _check(ws,c,r,empty,condition,message):
    _set(ws,c,r,f'IF({empty},"",IFERROR(IF({condition},"","{message}"),"{message}"))')

def _no_formulas(columns,r):return ','.join(f'NOT(_xlfn.ISFORMULA({_ref(c,r)}))' for c in columns)


def build_finance_view(wb, *, monthly_capacity=600):
    if not check_schema(wb):raise FinanceValidationError('Finance inputs must be created first')
    if not 1<=monthly_capacity<=12000:raise FinanceValidationError('Invalid monthly capacity')
    bind_names(wb); inp=wb[INPUT]
    ne=capacity(inp,'FF_Cash');ni=capacity(inp,'FF_Items');no=capacity(inp,'FF_Overrides');nr=2*ni
    # Preserve worksheet identities and order. Adding charts is not FF-2 scope.
    for title in (DATA,VIEW):
        if title not in wb.sheetnames:wb.create_sheet(title)
        ws=wb[title]
        if ws._charts or ws._images:raise FinanceValidationError('Unexpected object on finance-owned generated sheet')
        ws.delete_rows(1,ws.max_row)
    data,cash=wb[DATA],wb[VIEW]
    ids=_dr('A',nr); amounts=_dr('B',nr); dirs=_dr('C',nr); cats=_dr('E',nr)
    balances=_dr('F',nr); effective=_dr('J',nr); dates=_dr('K',nr);statuses=_dr('M',nr)
    oids=_rng('U',no);oa=_rng('V',no);od=_rng('W',no);oc=_rng('X',no);oreason=_rng('Y',no)
    eventids=_dr('U',ne);edates=_dr('V',ne);eamts=_dr('W',ne);edirs=_dr('X',ne);ecats=_dr('Y',ne);elinks=_dr('Z',ne)
    for i in range(ni):
        src=FIRST+i
        for part in range(2):
            r=FIRST+2*i+part
            base=_ref('J',src); gross=_ref('K',src);kind=_ref('Q',src);cert=f'{kind}="certificate"'
            active=f'AND({base}<>"",'+(cert if part else 'TRUE')+')'
            suffix=':retention' if part else ':net'
            _set(data,1,r,f'IF({active},IF({cert},{base}&"{suffix}",{base}),"")')
            withheld=f'ROUND(ROUND({gross},2)*FF_Retention_Rate,2)'
            amount=withheld if part else f'ROUND({gross},2)-{withheld}'
            _set(data,2,r,f'IF(A{r}="",0,IF({cert},{amount},ROUND({gross},2)))')
            _set(data,3,r,f'IF(A{r}="","",{_ref("L",src)})')
            exp=_ref('R',src) if part else _ref('M',src)
            cdate=exp if part else f'{exp}+FF_Credit_Days'
            _set(data,4,r,f'IF(A{r}="",0,IF({cert},IF({exp}="",0,{cdate}),IF({exp}="",0,{exp})))')
            _set(data,5,r,f'IF(A{r}="","",{_ref("N",src)})')
            _set(data,6,r,f'IF(A{r}="",0,IF({_ref("O",src)}="",0,{_ref("O",src)}))')
            ret='TRUE' if part else f'IF({cert},FALSE,{_true(_ref("P",src))})'
            _set(data,7,r,f'IF(A{r}="",FALSE,{ret})')
            _set(data,8,r,f'IF(A{r}="",0,SUMIFS({eamts},{elinks},A{r},{edates},"<="&FF_Actuals_Through,{edates},">"&F{r}))')
            _set(data,9,r,f'ROUND(B{r}-H{r},2)')
            match=f'MATCH(A{r},{oids},0)';has=f'COUNTIF({oids},A{r})>0'
            has_amount=f'SUMPRODUCT(1*({oids}=A{r}),1*ISNUMBER({oa}))>0'
            has_date=f'SUMPRODUCT(1*({oids}=A{r}),1*ISNUMBER({od}))>0'
            _set(data,10,r,f'IF(A{r}="",0,IF({has_amount},ROUND(INDEX({oa},{match}),2),I{r}))')
            _set(data,11,r,f'IF(A{r}="",0,IF({has},IF({_true(f"INDEX({oc},{match})")},0,IF({has_date},INDEX({od},{match}),D{r})),D{r}))')
            _set(data,12,r,f'IF(AND(A{r}<>"",{has}),INDEX({oreason},{match}),"")')
            _set(data,13,r,f'IF(A{r}="","",IF(J{r}=0,"settled",IF(K{r}=0,"undated",IF(K{r}<=FF_Actuals_Through,"overdue","scheduled"))))')
            condition=f'AND(COUNTIF({ids},A{r})=1,B{r}>=0,H{r}>=0,H{r}<=B{r},J{r}>=0,F{r}<=FF_Actuals_Through,D{r}<=2958465,K{r}<=2958465)'
            _check(data,14,r,f'A{r}=""',condition,'Reconcile item / settlement')
            _set(data,15,r,f'ROUND(J{r}-I{r},2)')
            _set(data,16,r,f'IF(M{r}="scheduled",K{r},0)')
            _set(data,17,r,f'IF(AND(M{r}="scheduled",C{r}="in"),J{r},0)')
            _set(data,18,r,f'IF(AND(M{r}="scheduled",C{r}="out"),J{r},0)')
        r=src
        b,a,dr,d,cat,bal,ret,kind,release=(_ref(c,r) for c in 'JKLMNOPQR')
        empty='AND('+','.join(f'{_ref(c,r)}=""' for c in 'JKLMNOPQRS')+')'
        condition=f'AND({_id(b)},{_money(a)},{_date(d,True)},{_date(bal,True)},{_date(release,True)},OR({dr}="in",{dr}="out"),OR({cat}="operating",{cat}="financing"),OR({ret}="",{_bool(ret)}),OR({kind}="manual",{kind}="certificate"),IF({kind}="certificate",AND({dr}="in",{cat}="operating",{d}<>"",{bal}="",NOT({_true(ret)})),{release}=""),IF({_true(ret)},AND({dr}="in",{cat}="operating"),TRUE),{_no_formulas("JKLMNOPQRS",r)})'
        _check(data,42,r,empty,condition,'Invalid item input')
    for i in range(ne):
        r=FIRST+i
        for c,s in enumerate('ABCDEFG',21):
            _set(data,c,r,f'IF({_ref(s,r)}="",'+('0' if s in 'BC' else '""')+f','+(f'ROUND({_ref(s,r)},2)' if s=='C' else _ref(s,r))+')')
        b,d,a,dr,cat,link,reason=(_ref(c,r) for c in 'ABCDEFG')
        empty='AND('+','.join(f'{_ref(c,r)}=""' for c in 'ABCDEFGH')+')'
        condition=f'AND({_id(b)},COUNTIF({eventids},U{r})=1,{_date(d)},{_money(a,True)},OR({dr}="in",{dr}="out"),OR({cat}="operating",{cat}="financing"),IF({a}<0,LEN(TRIM({reason}))>0,TRUE),IF({link}="",TRUE,COUNTIFS({ids},{link},{dirs},{dr},{cats},{cat})=1),{_no_formulas("ABCDEFGH",r)})'
        _check(data,28,r,empty,condition,'Invalid cash input / link')
        _set(data,29,r,f'IF(U{r}="","",IF(V{r}>FF_Actuals_Through,"future_actual",IF(AND(Z{r}="",W{r}<>0,NOT({_true("FF_Unlinked_Reconciled")})),"unlinked_actual","")))')
        _set(data,30,r,f'IF(AND(U{r}<>"",V{r}>=FF_Opening_Date,V{r}<=FF_Actuals_Through),V{r},0)')
        _set(data,31,r,f'IF(AND(AD{r}>0,X{r}="in"),W{r},0)')
        _set(data,32,r,f'IF(AND(AD{r}>0,X{r}="out"),W{r},0)')
    for i in range(no):
        r=FIRST+i;b,a,d,clear,reason=(_ref(c,r) for c in 'UVWXY')
        empty='AND('+','.join(f'{_ref(c,r)}=""' for c in 'UVWXYZ')+')'
        condition=f'AND({_id(b)},COUNTIF({oids},{b})=1,COUNTIF({ids},{b})=1,{_money(a,optional=True)},{_date(d,True)},OR({clear}="",{_bool(clear)}),NOT(AND({_true(clear)},{d}<>"")),LEN(TRIM({reason}))>0,{_no_formulas("UVWXYZ",r)})'
        _check(data,43,r,empty,condition,'Invalid override / reason')
    # Each dated event is evaluated at end of day. Duplicate same-day candidate
    # dates are harmless: balances are cumulative by date, not by row order.
    n=ne+nr+3;end=FIRST+n-1
    for i in range(n):
        r=FIRST+i
        if i<ne: expr=f'AD{FIRST+i}'
        elif i<ne+nr:expr=f'P{FIRST+i-ne}'
        else:expr=('FF_Opening_Date','FF_Actuals_Through','MAX(FF_Opening_Date,FF_Horizon_End)')[i-ne-nr]
        _set(data,34,r,f'IF({expr}>0,{expr},2958466)')
        small=f'SMALL($AH${FIRST}:$AH${end},{i+1})'
        _set(data,35,r,f'IF({small}=2958466,0,{small})')
        for c,amt,dt,count in ((36,'AE','AD',ne),(37,'AF','AD',ne),(38,'Q','P',nr),(39,'R','P',nr)):
            _set(data,c,r,f'IF(AI{r}=0,0,SUMIFS({_dr(amt,count)},{_dr(dt,count)},AI{r}))')
        parts=[f'SUMIFS({_dr(amt,count)},{_dr(dt,count)},">0",{_dr(dt,count)},"<="&AI{r})' for amt,dt,count in [('AE','AD',ne),('AF','AD',ne),('Q','P',nr),('R','P',nr)]]
        _set(data,40,r,f'IF(AI{r}=0,FF_Opening_Cash,ROUND(FF_Opening_Cash+{parts[0]}-{parts[1]}+{parts[2]}-{parts[3]},2))')
        _set(data,41,r,f'IF(AND(AI{r}>0,AN{r}=$B$6),AI{r},2958466)')
    settings=f'AND(ISTEXT(FF_Currency),LEN(TRIM(FF_Currency))>0,{_date("FF_Opening_Date")},{_date("FF_Actuals_Through")},FF_Opening_Date<=FF_Actuals_Through,{_money("FF_Opening_Cash",True)},ISNUMBER(FF_Credit_Days),FF_Credit_Days>=0,FF_Credit_Days<=36500,MOD(FF_Credit_Days,1)=0,ISNUMBER(FF_Retention_Rate),FF_Retention_Rate>=0,FF_Retention_Rate<=1,{_date("FF_Horizon_End",True)},'+','.join(_bool('FF_'+s) for s in ('Actuals_Complete','Receivables_Complete','Cash_Out_Complete','Unlinked_Reconciled'))+','+','.join(f'NOT(_xlfn.ISFORMULA({_ref("B",r)}))' for r in SETTINGS.values())+')'
    _set(data,2,2,f'IFERROR(IF({settings},0,1),1)')
    _set(data,2,3,f'IFERROR(B2+SUMPRODUCT(1*({_dr("N",nr)}<>""))+SUMPRODUCT(1*({_dr("AB",ne)}<>""))+SUMPRODUCT(1*({_dr("AP",ni)}<>""))+SUMPRODUCT(1*({_dr("AQ",no)}<>""))+IF(B10>{monthly_capacity},1,0),1)')
    _set(data,2,4,f'IFERROR(COUNTIF({statuses},"undated")+COUNTIF({statuses},"overdue")+SUMPRODUCT(1*({_dr("AC",ne)}<>""))+'+'+'.join(f'IF({_true("FF_"+s)},0,1)' for s in ('Actuals_Complete','Receivables_Complete','Cash_Out_Complete'))+',1)')
    _set(data,2,5,'IF(B3>0,"INVALID INPUT",IF(B4>0,"CONDITIONAL / INCOMPLETE","COMPLETE"))')
    _set(data,2,6,f'MIN(FF_Opening_Cash,$AN${FIRST}:$AN${end})')
    _set(data,2,7,f'IF(B6=FF_Opening_Cash,FF_Opening_Date,MIN($AO${FIRST}:$AO${end}))')
    _set(data,2,8,'MAX(0,-B6)')
    _set(data,2,9,f'MAX(FF_Opening_Date,FF_Actuals_Through,FF_Horizon_End,{_dr("P",nr)},{_dr("AD",ne)})')
    _set(data,2,10,'(YEAR(B9)-YEAR(FF_Opening_Date))*12+MONTH(B9)-MONTH(FF_Opening_Date)+1')
    cash['A2']='Cash flow'
    summaries=[('Forecast status',"Finance_Data!B5"),('Invalid input checks','Finance_Data!B3'),('Incomplete checks','Finance_Data!B4'),
        ('Peak funding requirement','Finance_Data!B8'),('Minimum cash date','Finance_Data!B7'),('Minimum cash','Finance_Data!B6'),
        ('Remaining receivable',f'SUMIFS(\'{DATA}\'!{effective},\'{DATA}\'!{dirs},"in",\'{DATA}\'!{cats},"operating")'),
        ('Remaining cash out',f'SUMIFS(\'{DATA}\'!{effective},\'{DATA}\'!{dirs},"out",\'{DATA}\'!{cats},"operating")'),
        ('Remaining retention',f'SUMIFS(\'{DATA}\'!{effective},\'{DATA}\'!{_dr("G",nr)},TRUE)'),
        ('Cash Out AC proxy',f'SUMIFS(\'{DATA}\'!{eamts},\'{DATA}\'!{edirs},"out",\'{DATA}\'!{ecats},"operating",\'{DATA}\'!{edates},"<="&FF_Actuals_Through)'),
        ('Currency','FF_Currency')]
    for r,(label,formula) in enumerate(summaries,4):
        cash.cell(r,1,label);_set(cash,2,r,formula if r in (4,5,6,14) else f'IF(Finance_Data!B3>0,NA(),{formula})')
    cash['D4']='Funding uses dated closing cash, including explicit financing.'
    cash['D5']='Undated and overdue balances stay in remaining totals.'
    cash['D6']='Conditional results exclude unresolved dated cash.'
    cash['D7']='Review Finance_Data item status and input checks before use.'
    cash['D8']='Monthly closing cash does not determine peak funding.'
    cash['D9']='Terminal cash is not profit.'
    for c,h in enumerate(('Month','Actual in','Actual out','Forecast in','Forecast out','Closing cash'),1):cash.cell(18,c,h)
    for i in range(monthly_capacity):
        r=19+i
        start=f'DATE(YEAR(FF_Opening_Date),MONTH(FF_Opening_Date)+{i},1)'
        _set(cash,1,r,f'IF({i}<Finance_Data!B10,{start},"")')
        for c,amt,dt,count in ((2,'AE','AD',ne),(3,'AF','AD',ne),(4,'Q','P',nr),(5,'R','P',nr)):
            _set(cash,c,r,f'IF(A{r}="","",IF(Finance_Data!B3>0,NA(),SUMIFS(\'{DATA}\'!{_dr(amt,count)},\'{DATA}\'!{_dr(dt,count)},">="&A{r},\'{DATA}\'!{_dr(dt,count)},"<"&DATE(YEAR(A{r}),MONTH(A{r})+1,1))))')
        prev='FF_Opening_Cash' if i==0 else f'F{r-1}'
        _set(cash,6,r,f'IF(A{r}="","",ROUND({prev}+B{r}-C{r}+D{r}-E{r},2))')
    headers=['ID','Original amount','Direction','Original date','Category','Balance date','Retention','Settled','Base remaining','Effective remaining','Effective date','Override reason','Status','Validation','Scenario adjustment','Scheduled date','Forecast in','Forecast out']
    for c,h in enumerate(headers,1):data.cell(FIRST-1,c,h)
    for c,h in enumerate(['Event ID','Date','Amount','Direction','Category','Link','Correction','Validation','Completeness','Series date','Actual in','Actual out'],21):data.cell(FIRST-1,c,h)
    for c,h in enumerate(['Candidate date','Sorted date','Actual in','Actual out','Forecast in','Forecast out','Closing cash','Min date candidate','Item input check','Override check'],34):data.cell(FIRST-1,c,h)
    for r,label in enumerate(['Settings invalid','Invalid checks','Incomplete checks','Status','Minimum cash','Minimum date','Peak funding','Last date','Months required'],2):data.cell(r,1,label)
    for ws in (data,cash):
        style_sheet(ws)
        for c in range(1,ws.max_column+1):ws.column_dimensions[col(c)].width=22
        ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=26
        hr=24 if ws==data else 18;ws.row_dimensions[hr].height=36
        for cell in ws[hr]:
            cell.font=Font(name=_FONT,size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True,vertical='center')
    for row in range(FIRST,FIRST+max(n,nr)):
        for c in (4,6,11,16,22,30,34,35,41):data.cell(row,c).number_format=DATE_FORMAT
        for c in (2,8,9,10,15,17,18,23,31,32,36,37,38,39,40):data.cell(row,c).number_format=MONEY_FORMAT
    for row in range(19,19+monthly_capacity):
        cash.cell(row,1).number_format='mmm-yyyy'
        for c in range(2,7):cash.cell(row,c).number_format=MONEY_FORMAT
    for r in range(7,14):cash.cell(r,2).number_format=DATE_FORMAT if r==8 else MONEY_FORMAT
    cash.print_options.horizontalCentered=True
    cash.print_area=f'A2:F{min(19+monthly_capacity,60)}'
