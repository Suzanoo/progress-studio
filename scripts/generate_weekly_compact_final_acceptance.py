"""Generate Final Presentation acceptance through the production Create boundary.
Synthetic input fixture only; all rendering is owned by production modules.
"""
from datetime import datetime, timedelta
from pathlib import Path
import argparse
import sys
import tempfile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from progress_studio.domain.activity import Activity
from progress_studio.infrastructure.excel.import_workbook_writer import ImportWorkbookWriter
from progress_studio.infrastructure.excel.schedule_workbook import transform_file
from progress_studio.infrastructure.excel.timescale_workbook import add_weekly_timescale
from progress_studio.infrastructure.excel.progress_workbook import prepare_progress_and_scurve
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.live_dashboard_workbook import build_live_dashboard
from progress_studio.services.monthly_main_service import MonthlyMainService
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.infrastructure.excel.calculation_policy import request_initial_manual_excel_recalculation
CUTOFF=datetime(2026,10,23)

def dt(s):
    return datetime.strptime(s, '%Y-%m-%d')


def dates(first, last):
    d = dt(first)
    out = []
    while d <= dt(last):
        out.append(d)
        d += timedelta(days=7)
    return out


def fixture():
    """Authored synthetic weekly observations; not a progress engine."""
    long_dates = dates('2026-09-04', '2026-12-18')
    plan = dict(zip(long_dates, [.075]*8 + [.05]*8))
    actual = dict(zip(long_dates[:9], [.05625]*8 + [.05]))
    completed = dates('2026-07-17', '2026-08-28')
    small_dates = dates('2026-10-02', '2027-01-22')
    small_plan = dict(zip(small_dates, [.0015]*4 + [.994/(len(small_dates)-4)]*(len(small_dates)-4)))
    small_actual = dict(zip(small_dates[:5], [.000625]*4 + [.0035]))
    gap = {d: .04 for d in long_dates[:9] if d != dt('2026-10-09')}
    gap_elsewhere = {d: .05 for d in long_dates[:9] if d != dt('2026-09-18')}
    future = dates('2026-11-06', '2027-02-12')
    zero = {d: 0.0 for d in dates('2026-10-02', '2026-10-16')}
    zero.update({d: .25 for d in dates('2026-11-06', '2026-11-27')})
    def activity(key, name, p, a, level=2, window='2026-10-02'):
        return dict(key=key, name=name, plan=p, actual=a, level=level,
                    first=min(p), last=max(p), summary=False, window=window)
    def wbs(key, name, first, last, level=1):
        return dict(key=key, name=name, first=dt(first), last=dt(last),
                    level=level, summary=True, window='2026-10-02')
    return [
        wbs('1', 'Civil works', '2026-07-17', '2026-12-18'),
        activity('CIV-010', 'Groundworks complete', dict.fromkeys(completed, 1/7), dict.fromkeys(completed, 1/7), window='2026-08-07'),
        wbs('1.1', 'Concrete structure', '2026-09-04', '2026-12-18', level=2),
        activity('CIV-020', 'Foundations', plan, actual, level=3),
        activity('CIV-030', 'Crane lift (short)', {dt('2026-10-16'): 1.0}, {dt('2026-10-16'): .5}, level=3),
        activity('CIV-040', 'Retaining walls', plan, dict.fromkeys(long_dates[:9], 0.0), level=3),
        wbs('2', 'Building services', '2026-09-04', '2027-02-12'),
        activity('MEP-010', 'First fix (window gap)', plan, gap),
        activity('MEP-020', 'Containment (earlier gap)', plan, gap_elsewhere),
        activity('MEP-030', 'Controls commissioning', small_plan, small_actual),
        activity('MEP-040', 'Plant testing (future)', dict.fromkeys(future, 1/len(future)), {}, window='2026-12-04'),
        activity('MEP-050', 'Fire systems (no Actual)', plan, {}),
        activity('MEP-060', 'BMS release (zero to cutoff)', zero, {}),
    ]



def build_acceptance(path):
    specs=fixture()
    long_dates=dates('2024-01-12','2027-02-12')
    specs.append(dict(key='LONG-010',name='Multi-year horizon',first=long_dates[0],last=long_dates[-1],level=1,summary=False,plan=dict.fromkeys(long_dates,1/len(long_dates)),actual={d:.7/len(long_dates) for d in long_dates if d<=CUTOFF}))
    with tempfile.TemporaryDirectory() as td:
        folder=Path(td)
        rows=[Activity(i,i,i,s['key'],s['name'],s['key'] if s['summary'] else '',s['level'],s['summary'],s['first'],s['last'],None,None,None,None,0,100000.) for i,s in enumerate(specs,1)]
        imported,paired,dated,ready=[folder/(n+'.xlsx') for n in ['imported','paired','dated','ready']]
        ImportWorkbookWriter().write(imported,Path('SYNTHETIC'),'Weekly Compact Final Presentation',rows)
        transform_file(imported,paired,'main',repeat_description=False)
        add_weekly_timescale(paired,dated,'main',cutoff_day=5,margin_weeks=3)
        w=load_workbook(dated);ds=main_dataset_from_workbook(w);main=w['main'];by_id={s['key']:s for s in specs if not s['summary']}
        for row in ds.activities:
            spec=by_id[row.activity_id]
            main.cell(row.row_number,ds.header_column('amount'),100000.)
            for p in ds.periods:
                main.cell(row.row_number,p.column).value=spec['plan'].get(p.reporting_date)
                main.cell(row.row_number+1,p.column).value=spec['actual'].get(p.reporting_date)
        prepare_progress_and_scurve(w,main);ds=main_dataset_from_workbook(w)
        build_live_dashboard(w,ds,project_name='Weekly Compact Final Presentation')
        w.save(ready);w.close()
        path.parent.mkdir(parents=True,exist_ok=True)
        MonthlyMainService().build(ready,path)
        # Set one ordinary user input for repeatable inspection, then exercise
        # the production preservation/finalization path without proof adapters.
        w=load_workbook(path);sheet,cell=next(w.defined_names['PS_WEEKLY_OVERLAY_CUTOFF'].destinations)
        w[sheet][cell]=CUTOFF
        finalize_workbook(w,mode='live',include_guide=True)
        request_initial_manual_excel_recalculation(w)
        w.active=w.sheetnames.index('Weekly Compact')
        w['Weekly Compact'].sheet_view.pane.topLeftCell='FC5'
        w.save(path);w.close()
    return path

if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,default=Path('WEEKLY_COMPACT_FINAL_PRESENTATION_ACCEPTANCE.xlsx'))
    print(build_acceptance(parser.parse_args().output))
