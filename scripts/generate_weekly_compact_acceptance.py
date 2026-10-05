"""Deterministic synthetic V1-A inspection workbook, using production renderers.

Run from repository root: python scripts/generate_weekly_compact_acceptance.py
No Git operations; no OOXML patching. Desktop acceptance remains a human gate.
"""
from __future__ import annotations
import argparse
import math
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from progress_studio.domain.activity import Activity
from progress_studio.infrastructure.excel.import_workbook_writer import ImportWorkbookWriter
from progress_studio.infrastructure.excel.schedule_workbook import transform_file
from progress_studio.infrastructure.excel.timescale_workbook import add_weekly_timescale
from progress_studio.infrastructure.excel.progress_workbook import prepare_progress_and_scurve
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.live_dashboard_workbook import build_live_dashboard
from progress_studio.infrastructure.excel.live_monthly_workbook import build_live_monthly_view
from progress_studio.services.monthly_cache_deriver import MonthlyCacheDeriver
from progress_studio.infrastructure.excel.traditional_overlay_workbook import build_traditional_overlays
from progress_studio.infrastructure.excel.weekly_compact_workbook import build_weekly_compact
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.infrastructure.excel.calculation_policy import request_initial_manual_excel_recalculation

START = datetime(2022, 1, 14)
FINISH = datetime(2027, 12, 17)
CUTOFF = datetime(2026, 10, 23)


def build_source(folder):
    rows = []
    names = ["Site preparation", "Foundations", "Substructure", "Superstructure",
             "Roof structure", "Building envelope", "MEP first fix", "Partitions",
             "Finishes", "MEP final fix", "External works", "Testing and handover"]
    for i, name in enumerate(["Synthetic long-project schedule"] + names):
        first = START if i == 0 else START + timedelta(weeks=max(0, (i-1)*16))
        last = FINISH if i in (0, 12) else min(FINISH, first + timedelta(weeks=132))
        rows.append(Activity(i, i, i, f"DEMO-{i:02}", name, "1", 1 if i == 0 else 2,
                             i == 0, first, last, None, None, None, None, 0, 100.0))
    imported, paired, dated = [folder / f"{name}.xlsx" for name in ("imported", "paired", "dated")]
    ImportWorkbookWriter().write(imported, Path("SYNTHETIC-NOT-XML"), "Weekly Compact V1-A (synthetic)", rows)
    transform_file(imported, paired, "main", repeat_description=False)
    add_weekly_timescale(paired, dated, "main", cutoff_day=5, margin_weeks=3)
    wb = load_workbook(dated)
    ws = wb["main"]
    ds = main_dataset_from_workbook(wb)
    h = dict(ds.headers)
    periods = ds.periods
    for row in ds.activities:
        eligible = [p for p in periods if row.plan_start <= p.reporting_date <= row.plan_finish]
        weights = [math.exp(-((i/max(1,len(eligible)-1)-.5)/.28)**2) for i in range(len(eligible))]
        total = sum(weights)
        for i, (p, weight) in enumerate(zip(eligible, weights)):
            ws.cell(row.row_number, p.column, weight/total)
            if p.reporting_date <= CUTOFF + timedelta(days=7):
                # Synthetic observed weekly increments, with deliberate blanks.
                value = weight/total * (.78 + .08*math.sin(i*.53))
                ws.cell(row.row_number+1, p.column, None if i in (9, 10) else value)
    prepare_progress_and_scurve(wb, ws)
    return wb


def build_acceptance(path):
    with tempfile.TemporaryDirectory(prefix="weekly-compact-") as tmp:
        wb = build_source(Path(tmp))
        ds = main_dataset_from_workbook(wb)
        build_live_monthly_view(wb, ds, MonthlyCacheDeriver().derive(ds))
        build_live_dashboard(wb, ds, project_name="Weekly Compact V1-A (synthetic)")
        build_traditional_overlays(wb, ds)
        name = wb.defined_names["PS_WEEKLY_OVERLAY_CUTOFF"]
        sheet, address = next(name.destinations)
        wb[sheet][address] = CUTOFF
        build_weekly_compact(wb, ds)
        finalize_workbook(wb, mode="live", include_guide=True)
        request_initial_manual_excel_recalculation(wb)
        wb.active = wb.sheetnames.index("Weekly Compact")
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        wb.close()
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("WEEKLY_COMPACT_V1A_ACCEPTANCE.xlsx"))
    args = parser.parse_args()
    print(build_acceptance(args.output))
