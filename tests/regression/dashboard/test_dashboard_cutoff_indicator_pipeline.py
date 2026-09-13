"""Cutoff contracts through real rebuild/overlay output, not Excel rendering."""
from datetime import datetime, timedelta
import math
import posixpath
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from openpyxl.utils.datetime import to_excel

from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from tests.integration.rebuild.test_progress_rebuild_engine import _full_rebuild_fixture
from tests.regression.dashboard.test_live_marker_alignment import _evaluate_marker

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      's': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def _related(package, part, relation_id):
    folder, name = posixpath.split(part)
    relationships = ET.fromstring(package.read(f'{folder}/_rels/{name}.rels'))
    target = next(r.get('Target') for r in relationships if r.get('Id') == relation_id)
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(folder, target))


def chart_for_sheet(package, sheet_name):
    """Follow worksheet/drawing relationships; never assume chart package order."""
    book = ET.fromstring(package.read('xl/workbook.xml'))
    sheet = next(s for s in book.findall('s:sheets/s:sheet', NS) if s.get('name') == sheet_name)
    part = _related(package, 'xl/workbook.xml', sheet.get('{'+NS['r']+'}id'))
    drawing = ET.fromstring(package.read(part)).find('s:drawing', NS)
    part = _related(package, part, drawing.get('{'+NS['r']+'}id'))
    charts = ET.fromstring(package.read(part)).findall('.//c:chart', NS)
    assert len(charts) == 1
    part = _related(package, part, charts[0].get('{'+NS['r']+'}id'))
    return ET.fromstring(package.read(part))


def assert_date_label(series, prefix='Cutoff'):
    label = series.find('c:dLbls', NS)
    assert label.find('c:showCatName', NS).get('val') == '1'
    assert label.find('c:showSerName', NS).get('val') == '0'
    assert label.find('c:showVal', NS).get('val') == '0'
    assert label.find('c:dLblPos', NS).get('val') == 't'
    assert label.find('c:separator', NS) is None
    assert label.find('c:numFmt', NS).get('formatCode') == f'"{prefix} "dd/mm/yyyy'
    assert label.find('c:spPr/a:solidFill/a:srgbClr', NS).get('val') == 'FCE4D6'
    assert label.find('c:spPr/a:ln/a:solidFill/a:srgbClr', NS).get('val') == 'C00000'
    run = label.find('c:txPr/a:p/a:pPr/a:defRPr', NS)
    assert run.get('b') == '1' and run.get('sz') == '1000'
    assert run.find('a:solidFill/a:srgbClr', NS).get('val') == 'C00000'
    assert label.findall('c:dLbl', NS) == []  # no persisted per-point label list


def _source(path):
    _full_rebuild_fixture(path)
    wb = load_workbook(path)
    ws = wb['main']
    # Three months and multiple weeks exercise different view lengths.
    for row in (5, 6, 9, 10, 11, 12):
        ws.cell(row, 10, datetime(2026, 5, 31))
    for index in range(13):
        col = 12 + index
        ws.cell(3, col, f'W{index + 1}')
        ws.cell(4, col, datetime(2026, 3, 6) + timedelta(days=7 * index))
        for row in (5, 9, 11):
            ws.cell(row, col, 1 / 13)
        for row in (6, 10, 12):
            ws.cell(row, col, .02 if index < 6 else 0)
        ws.cell(7, col, (index + 1) / 13)
        ws.cell(8, col, min(index + 1, 6) * .02)
    wb.save(path)
    wb.close()
    return path


def _cell_value(value):
    # Excel round-trip normalizes date -> datetime and empty string -> None.
    if isinstance(value, float):
        return round(value, 12)
    return value.date() if isinstance(value, datetime) else None if value == '' else value


@pytest.mark.parametrize('mode', ['snapshot', 'live'])
def test_dashboard_indicator_survives_actual_pipeline_and_reopen(tmp_path, monkeypatch, mode):
    from progress_studio.infrastructure.excel import dashboard_workbook as normal
    from progress_studio.infrastructure.excel import live_dashboard_workbook as live
    captured = {}
    original = normal._add_dashboard_cutoff_indicator

    def capture(chart, data, cats, last_row):
        # Capture accepted curves/CP-1 markers BEFORE the new helper is added.
        captured['series'] = [ET.tostring(s.to_tree()) for s in chart.series]
        captured['cells'] = {(r, c): data.cell(r, c).value
                             for r in range(1, last_row + 1) for c in range(1, 16)}
        captured['axes'] = (ET.tostring(chart.x_axis.to_tree()), ET.tostring(chart.y_axis.to_tree()))
        assert all(data.cell(r, 28).value is None for r in range(1, last_row + 1))
        original(chart, data, cats, last_row)
        assert [ET.tostring(s.to_tree()) for s in chart.series[:-1]] == captured['series']
        captured['ab'] = [data.cell(r, 28).value for r in range(1, last_row + 1)]

    monkeypatch.setattr(normal, '_add_dashboard_cutoff_indicator', capture)
    monkeypatch.setattr(live, '_add_dashboard_cutoff_indicator', capture)
    source = _source(tmp_path / 'source.xlsx')
    output = tmp_path / 'output.xlsx'
    engine = WorkbookRebuildEngine()
    method = engine.rebuild_live_progress if mode == 'live' else engine.rebuild_progress
    method(source, output, project_name='Cutoff pipeline')
    for iteration in range(2):
        wb = load_workbook(output)
        data = wb['Dashboard_Data']
        chart = wb['Dashboard']._charts[0]
        assert len(chart.series) == (5 if mode == 'live' else 3)
        last_row = range_to_tuple(chart.series[-1].val.numRef.f)[1][3]
        assert [data.cell(r, 28).value for r in range(1, last_row + 1)] == captured['ab']
        assert {(r, c): _cell_value(data.cell(r, c).value) for r, c in captured['cells']} == {
            key: _cell_value(value) for key, value in captured['cells'].items()}
        assert data['P1'].value == 'Weekly Actual Visible'
        assert data['AA1'].value == 'Monthly Overlay Cutoff'
        assert data['P2'].value != data['AB2'].value
        assert len(wb['main']._charts) == len(wb['main_monthly']._charts) == 1
        assert wb['Payment']['A1'].value == 'KEEP_PAYMENT'
        assert wb['Payment Input']['A1'].value == 'KEEP_PAYMENT_INPUT'
        assert wb.calculation.calcMode == 'manual' and wb.calculation.calcOnSave
        assert wb['Dashboard'].protection.sheet
        assert (ET.tostring(chart.x_axis.to_tree()), ET.tostring(chart.y_axis.to_tree())) == captured['axes']
        # Formulas remain live; supplied numeric reporting dates model Excel's
        # calculated G column, not an assertion that openpyxl calculates Excel.
        weekly = [to_excel(data.cell(r, 10).value) if data.cell(r, 10).value else ''
                  for r in range(2, last_row + 1)]
        monthly = [to_excel(data.cell(r, 11).value) if data.cell(r, 11).value else ''
                   for r in range(2, last_row + 1)]
        for view, dates in [('Weekly', weekly), ('Monthly', monthly)]:
            valid = [d for d in dates if d != '']
            assert len(valid) >= 3
            for row in range(2, last_row + 1):
                assert _evaluate_marker(data.cell(row, 7).value, {
                    'Dashboard!G5': view, f'A{row}': weekly[row - 2],
                    f'D{row}': monthly[row - 2]}) == dates[row - 2]
            cutoffs = (valid[0] - 1, valid[0], valid[0] + 1,
                       valid[len(valid) // 2], valid[len(valid) // 2] + 1,
                       valid[-1], valid[-1] + 1)
            for cutoff in cutoffs:
                wb['Dashboard']['G5'] = view
                wb['Dashboard']['K5'] = cutoff
                hits = []
                for row in range(2, last_row + 1):
                    current = dates[row - 2]
                    following = dates[row - 1] if row - 1 < len(dates) else ''
                    # The evaluator eagerly evaluates branches. Use numeric 0
                    # for empty cells in comparisons, as Excel does, while
                    # retaining the formula's explicit blank guards separately.
                    formula = data.cell(row, 28).value
                    assert f'IF(G{row}="",NA(),' in formula
                    formula = formula.replace(f'G{row}=""', '1=1' if current == '' else '1=0')
                    formula = formula.replace(f'G{row + 1}=""', '1=1' if following == '' else '1=0')
                    value = _evaluate_marker(formula, {f'G{row}': current or 0,
                        f'G{row + 1}': following or 0, 'Dashboard!K5': cutoff})
                    if not math.isnan(value):
                        hits.append((current, value))
                eligible = [d for d in dates if d != '' and d <= cutoff]
                assert hits == ([(max(eligible), 1)] if eligible else [])
        with ZipFile(output) as package:
            root = chart_for_sheet(package, 'Dashboard')
            series = root.findall('.//c:lineChart/c:ser', NS)
            cutoff = series[-1]
            assert cutoff.find('c:cat/c:numRef/c:f', NS).text == f"'Dashboard_Data'!$G$2:$G${last_row}"
            assert cutoff.find('c:val/c:numRef/c:f', NS).text == f"'Dashboard_Data'!$AB$2:$AB${last_row}"
            for s in series[:-1]:
                assert s.find('c:dLbls', NS) is None
                assert s.find('c:cat/c:numRef/c:f', NS).text == cutoff.find('c:cat/c:numRef/c:f', NS).text
            assert_date_label(cutoff)
            assert cutoff.find('c:tx/c:v', NS).text == 'Cutoff'
            error = cutoff.find('c:errBars', NS)
            assert error.find('c:errDir', NS).get('val') == 'y'
            assert error.find('c:errBarType', NS).get('val') == 'minus'
            assert error.find('c:val', NS).get('val') == '1'
            assert error.find('c:noEndCap', NS).get('val') == '1'
            assert error.find('c:spPr/a:ln/a:prstDash', NS).get('val') == 'dash'
            assert error.find('c:spPr/a:ln/a:solidFill/a:srgbClr', NS).get('val') == 'C00000'
            hidden = [(int(e.find('c:idx', NS).get('val')), e.find('c:delete', NS).get('val'))
                      for e in root.findall('.//c:legend/c:legendEntry', NS)]
            assert hidden == ([(2, '1'), (3, '1'), (4, '1')] if mode == 'live' else [(2, '1')])
        if iteration == 0:
            output = tmp_path / 'reopened.xlsx'
            wb.save(output)
        wb.close()
