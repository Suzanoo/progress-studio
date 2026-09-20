"""Totals use rendered monthly periods, not copied weekly offsets."""
from datetime import datetime
import re
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from tests.regression.dashboard.test_monthly_main_view import _workbook
from tests.fixtures.amount_xml import amount_xml
from progress_studio.infrastructure.excel.monthly_main_workbook import build_monthly_main_view
from progress_studio.app.desktop import DesktopRunner, DesktopRunOptions
from progress_studio.infrastructure.schedule_xml import NormalizedScheduleXmlReader
from progress_studio.services.rebuild_service import WorkbookRebuildEngine

@pytest.mark.parametrize('values,expected', [([1,None],1),([.25,.75],1),([None,''],''),([0,None],0)])
def test_early_late_blank_and_zero(values,expected):
    wb=_workbook(); ws=wb['main']
    for col in range(18,25):
        ws.cell(3,col,f'W{col-17}'); ws.cell(4,col,datetime(2026,col-17,23))
    for row in range(5,11):
        ws.cell(row,13,f'=IF(COUNT(W{row}:CK{row})=0,"",SUM(W{row}:CK{row}))')
    build_monthly_main_view(wb); monthly=wb['main_monthly']
    for row in range(5,11):  # Project/WBS/Activity P and A
        formula=monthly.cell(row,13).value
        assert formula==f'=IF(COUNT(R{row}:X{row})=0,"",SUM(R{row}:X{row}))'
        for col in range(18,25): monthly.cell(row,col).value=None
        monthly.cell(row,18).value=values[0]; monthly.cell(row,24).value=values[1]
        span=re.search(r'COUNT\(([^)]+)\)',formula)[1]
        nums=[c.value for cells in monthly[span] for c in cells if isinstance(c.value,(int,float)) and not isinstance(c.value,bool)]
        # Restricted COUNT/SUM numeric semantics, not a native Excel evaluator.
        assert (sum(nums) if nums else '')==expected
    assert ws['M9'].value=='=IF(COUNT(W9:CK9)=0,"",SUM(W9:CK9))'

def test_excludes_display_only_months():
    wb=_workbook(); wb['main']['R3']='X'
    build_monthly_main_view(wb)
    assert wb['main_monthly']['M9'].value=='=IF(COUNT(S9:S9)=0,"",SUM(S9:S9))'

def test_snapshot_keeps_cached_totals():
    wb=_workbook(); wb['main']['M9']='=IF(COUNT(R9:U9)=0,"",SUM(R9:U9))'
    values=_workbook()['main']; values['M9']=1
    build_monthly_main_view(wb,snapshot=True,value_source=values)
    assert wb['main_monthly']['M9'].value==1

def assert_ranges(wb):
    for name,prefix in [('main','W'),('main_monthly','M')]:
        ws=wb[name]
        cols=[c for c in range(1,ws.max_column+1) if re.fullmatch(prefix+r'\d+',str(ws.cell(3,c).value or ''))]
        headers={str(ws.cell(4,c).value).lower():c for c in range(1,ws.max_column+1)}
        checked=0
        for row in range(5,ws.max_row+1):
            if ws.cell(row,headers['p/a']).value not in ('P','A'): continue
            if str(ws.cell(row,headers['row type']).value).lower()=='s-curve': continue
            formula=ws.cell(row,headers['% complete']).value
            if not isinstance(formula,str) or not formula.startswith('='): continue
            span=f'{get_column_letter(cols[0])}{row}:{get_column_letter(cols[-1])}{row}'
            assert formula==f'=IF(COUNT({span})=0,"",SUM({span}))'
            checked+=1
        assert checked>0

@pytest.mark.parametrize('source',['p6','msp'])
@pytest.mark.parametrize('basis',['equal','duration','amount'])
def test_create_and_live_rebuild_ranges(tmp_path,monkeypatch,source,basis):
    xml=amount_xml(tmp_path,source)
    field=NormalizedScheduleXmlReader().read_with_amount_fields(xml)[2][0]
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path',lambda:tmp_path/'desktop')
    result=DesktopRunner().run(DesktopRunOptions(xml,'5',weight_basis=basis,amount_field=field.identity if basis=='amount' else None,distribution_method='flat'))
    wb=load_workbook(result.output_workbook); assert_ranges(wb); wb.close()
    output=tmp_path/'live.xlsx'
    WorkbookRebuildEngine().rebuild_live_progress(result.output_workbook,output)
    wb=load_workbook(output); assert_ranges(wb); wb.close()
