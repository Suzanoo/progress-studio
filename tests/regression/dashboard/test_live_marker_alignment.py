"""CP-1 package/formula checks, not desktop Excel visual acceptance."""
from datetime import datetime
import math
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer

from progress_studio.infrastructure.excel.live_dashboard_workbook import _build_live_data_sheet, build_live_dashboard
from progress_studio.infrastructure.excel.rebuild_workbook_reader import RebuildWorkbookReader
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from tests.regression.dashboard.test_progress_curve_contract import _fixture
from tests.regression.test_df1_dashboard_actual_cutoff_contract import _dataset, _workbook


def _evaluate_marker(formula, cells):
    """Scalar marker evaluator over supplied calculated G/H/L and weekly history.

    This deliberately does not evaluate the upstream Excel workbook or rendering.
    """
    def value(ref):
        return cells[ref.replace('$', '')]

    def countifs(dates, criterion, actual, nonblank):
        assert criterion.startswith('<=') and nonblank == '<>'
        limit = float(criterion[2:])
        return sum(d <= limit and a not in (None, '') for d, a in zip(dates, actual))

    formula = re.sub(r'"<="&([A-Za-z0-9!$]+)', lambda m: '"<=' + str(value(m[1])) + '"', formula)
    parts = []
    for token in Tokenizer(formula).items:
        if token.type == 'OPERAND' and token.subtype == 'RANGE':
            parts.append(repr(value(token.value)))
        elif token.type == 'OPERATOR-INFIX':
            parts.append({'=': '==', '<>': '!='}.get(token.value, token.value))
        elif token.type == 'FUNC' and token.subtype == 'OPEN':
            parts.append(token.value.lower())
        else:
            parts.append(token.value)
    functions = {'if_': lambda c, y, n: y if c else n,
                 'or_': lambda *args: any(args), 'and_': lambda *args: all(args),
                 'abs': abs, 'na': lambda: float('nan'),
                 'iferror': lambda result, fallback: result, 'countifs': countifs}
    expression = ''.join(parts)
    for name in ('if', 'or', 'and'):
        expression = re.sub(r'\b' + name + r'\(', name + '_(', expression)
    return eval(expression, {'__builtins__': {}}, functions)


@pytest.mark.parametrize('view,selected', [('Weekly', 2), ('Weekly', 6), ('Monthly', 2), ('Monthly', 3)])
@pytest.mark.parametrize('history', ['present', 'missing', 'zero', 'coincident'])
def test_markers_select_only_cutoff_and_distinguish_missing_from_zero(view, selected, history):
    wb = _workbook()
    _build_live_data_sheet(wb, _dataset(), object())
    data = wb['Dashboard_Data']
    dates = list(range(1, 7)) if view == 'Weekly' else [4, 6, '', '', '', '']
    cutoff = dates[selected - 2]
    for row, reporting in enumerate(dates, 2):
        source_cells = {'Dashboard!G5': view, f'A{row}': row - 1,
                        f'D{row}': [4, 6, '', '', '', ''][row - 2]}
        assert _evaluate_marker(data[f'G{row}'].value, source_cells) == reporting
    actual = [None, .2, '', '', '', '']
    if history == 'missing':
        actual = [None] * 6
    elif history == 'zero':
        actual = [None, 0, '', '', '', '']
    elif history == 'coincident':
        actual = [.8] * 6
    points = {14: [], 15: []}
    for row, reporting in enumerate(dates, 2):
        previous = [a for d, a in zip(range(1, 7), actual)
                    if reporting != '' and d <= reporting and a not in (None, '')]
        cells = {f'G{row}': reporting, f'H{row}': .8,
                 f'L{row}': previous[-1] if previous else 0,
                 'Dashboard!K5': cutoff, 'J2:J7': list(range(1, 7)), 'C2:C7': actual}
        for column in points:
            formula = data.cell(row, column).value
            assert formula and 'Dashboard!$K$5' in formula
            assert f'G{row}<>Dashboard!$K$5' in formula
            result = _evaluate_marker(formula, cells)
            if not math.isnan(result):
                points[column].append((row, result))
    assert points[14] == [(selected, .8)]
    expected = [] if history in ('missing', 'coincident') or cutoff < 2 else [(selected, 0 if history == 'zero' else .2)]
    assert points[15] == expected
    assert 'Dashboard!$G$5' in data['G2'].value
    assert 'IF(D7="","",D7)' in data['G7'].value
    wb.close()


def test_marker_ranges_and_axes_survive_save_and_round_trip(tmp_path):
    source = _fixture(tmp_path / 'source.xlsx')
    dataset = RebuildWorkbookReader().read_main_dataset(source)
    wb = load_workbook(source)
    build_live_dashboard(wb, dataset, cutoff=datetime(2026, 7, 24))
    for iteration in range(2):
        path = tmp_path / f'roundtrip-{iteration}.xlsx'
        finalize_workbook(wb, mode="live", include_guide=False)
        wb.save(path)
        wb.close()
        with ZipFile(path) as package:
            ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
            chart = ET.fromstring(package.read('xl/charts/chart1.xml'))
            series = chart.findall('.//c:lineChart/c:ser', ns)
            assert len(series) == 5
            for index, col in ((2, 'N'), (3, 'O')):
                category = series[index].find('c:cat/c:numRef/c:f', ns)
                assert category is not None
                assert category.text == "'Dashboard_Data'!$G$2:$G$6"
                assert series[index].find('c:val/c:numRef/c:f', ns).text == f"'Dashboard_Data'!${col}$2:${col}$6"
                assert series[index].find('c:dLbls', ns) is None
            axes = [chart.find('.//c:dateAx', ns), chart.find('.//c:valAx', ns)]
            assert [(a.find('c:axId', ns).get('val'), a.find('c:crossAx', ns).get('val')) for a in axes] == [('10', '100'), ('100', '10')]
        wb = load_workbook(path)
        for series in wb['Dashboard']._charts[0].series[2:4]:
            assert series.dLbls is None
    wb.close()


@pytest.mark.parametrize('actual', [None, '', 0, .2])
def test_actual_source_guard_preserves_blank_and_recorded_zero(actual):
    wb = _workbook()
    _build_live_data_sheet(wb, _dataset(), object())
    # Excel compares an empty cell to an empty string in the IF blank guard.
    value = '' if actual is None else actual
    result = _evaluate_marker(wb['Dashboard_Data']['C2'].value, {"'progress'!C2": value})
    assert result == value
    wb.close()


def test_cutoff_outside_selected_categories_does_not_fall_back_to_first_point():
    wb = _workbook()
    _build_live_data_sheet(wb, _dataset(), object())
    for row in range(2, 8):
        cells = {f'G{row}': row - 1, f'H{row}': .8, f'L{row}': .2,
                 'Dashboard!K5': 99, 'J2:J7': list(range(1, 7)), 'C2:C7': [.2] * 6}
        for column in (14, 15):
            assert math.isnan(_evaluate_marker(wb['Dashboard_Data'].cell(row, column).value, cells))
    wb.close()
