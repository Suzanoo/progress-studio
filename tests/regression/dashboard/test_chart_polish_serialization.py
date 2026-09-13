"""CP-2 serialized presentation contract; Excel visual acceptance is separate."""
from datetime import datetime
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_to_tuple

from progress_studio.infrastructure.excel.earned_value_workbook import render_earned_value_sheet
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.infrastructure.excel.live_dashboard_workbook import build_live_dashboard
from progress_studio.infrastructure.excel.rebuild_workbook_reader import RebuildWorkbookReader
from tests.regression.dashboard.test_progress_curve_contract import _fixture
from tests.unit.earn_value.test_earned_value_workbook import _result

NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def _title(series, wb):
    literal = series.find('c:tx/c:v', NS)
    if literal is not None:
        return literal.text
    reference = series.find('c:tx/c:strRef/c:f', NS).text
    name, (col, row, _, _) = range_to_tuple(reference)
    return wb[name].cell(row, col).value


@pytest.mark.parametrize('kind', ['Dashboard', 'Earned Value'])
def test_business_legend_and_marker_style_survive_final_save_and_reopen(tmp_path, kind):
    if kind == 'Dashboard':
        source = _fixture(tmp_path / 'input.xlsx')
        wb = load_workbook(source)
        build_live_dashboard(wb, RebuildWorkbookReader().read_main_dataset(source), cutoff=datetime(2026, 7, 24))
        titles = ['Plan', 'Actual', 'Cutoff Plan Marker', 'Cutoff Actual Marker']
        hidden = [2, 3]
    else:
        wb = Workbook()
        render_earned_value_sheet(wb, _result())
        assert wb[kind]['A3'].value == 'PROJECT PERFORMANCE'
        assert wb[kind]['A10'].value == 'PERFORMANCE CURVE — PV VS EV'
        titles = ['PV', 'EV', 'Status Date', 'PV @ Status Date', 'EV @ Status Date']
        hidden = [3, 4]
    initial_images = len(wb[kind]._images)
    original_refs = [[n.text for n in s.to_tree().iter() if n.tag.endswith('}f') or n.tag == 'f']
                     for s in wb[kind]._charts[0].series]
    for iteration in range(2):
        finalize_workbook(wb, include_guide=False)
        path = tmp_path / f'{iteration}.xlsx'
        wb.save(path)
        with ZipFile(path) as package:
            root = ET.fromstring(package.read('xl/charts/chart1.xml'))
            series = root.findall('.//c:lineChart/c:ser', NS)
            # Deletion indices are valid only with this verified serialized order.
            assert [_title(s, wb) for s in series] == titles
            assert [int(s.find('c:idx', NS).get('val')) for s in series] == list(range(len(titles)))
            entries = root.findall('.//c:legend/c:legendEntry', NS)
            assert [(int(e.find('c:idx', NS).get('val')), e.find('c:delete', NS).get('val'))
                    for e in entries] == [(i, '1') for i in hidden]
            for index, color in zip(hidden, ['2F75B5', '70AD47']):
                marker = series[index].find('c:marker', NS)
                assert marker.find('c:size', NS).get('val') == '7'
                assert marker.find('c:spPr/a:solidFill/a:srgbClr', NS).get('val') == color
                assert marker.find('c:spPr/a:ln/a:solidFill/a:srgbClr', NS).get('val') == color
                assert series[index].find('c:dLbls', NS) is None
            for axis, axis_id, cross_id in [('dateAx', '10', '100'), ('valAx', '100', '10')]:
                node = root.find('.//c:' + axis, NS)
                assert node.find('c:axId', NS).get('val') == axis_id
                assert node.find('c:crossAx', NS).get('val') == cross_id
            if kind == 'Earned Value':
                status = series[2]
                assert status.find('c:errBars/c:spPr/a:ln/a:solidFill/a:srgbClr', NS).get('val') == 'C00000'
                assert status.find('c:dLbls/c:showCatName', NS).get('val') == '1'
        wb.close()
        wb = load_workbook(path)
        assert len(wb[kind]._charts) == 1
        assert len(wb[kind]._images) == initial_images
        refs = [[n.text for n in s.to_tree().iter() if n.tag.endswith('}f') or n.tag == 'f']
                for s in wb[kind]._charts[0].series]
        assert refs == original_refs
    wb.close()
