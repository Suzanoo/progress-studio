from datetime import date
from pathlib import Path
import pytest
from openpyxl import load_workbook
from progress_studio.services.financial_forecast_workbook_service import FinancialForecastWorkbookService
from progress_studio.infrastructure.excel.finance_input_workbook import read_finance_inputs
from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from tests.fixtures.finance_workbook_probe import create_seed, package_issues


@pytest.fixture(scope='module')
def financed_seed(tmp_path_factory):
    directory=tmp_path_factory.mktemp('finance-lifecycle')
    source=create_seed(directory/'seed')
    # FF-0 fixture creates an unextended Progress/Payment/EV workbook.
    if isinstance(source,tuple):source=source[0]
    out=directory/'finance.xlsx'
    FinancialForecastWorkbookService().generate(source,out,opening_date=date(2026,1,1),actuals_through=date(2026,1,31),currency='THB',capacity_rows=4)
    w=load_workbook(out);s=w['Finance Input'];s['J25']='remain';s['K25']=100;s['L25']='in';s['M25']=date(2026,3,2);s['N25']='operating';s['Q25']='manual';s['S25']='Keep note'
    s['U28']='remain';s['V28']=0;s['Y28']='Cancelled';w.save(out);w.close()
    return out


@pytest.mark.parametrize('operation',['rebuild_progress','rebuild_live_progress','rebuild_payment','rebuild_live_payment'])
def test_finance_survives_progress_and_payment(financed_seed,tmp_path,operation):
    source=load_workbook(financed_seed);expected=read_finance_inputs(source);source.close()
    out=tmp_path/'rebuilt.xlsx';getattr(WorkbookRebuildEngine(),operation)(financed_seed,out)
    w=load_workbook(out)
    assert read_finance_inputs(w)==expected
    assert w['Finance Input']['S25'].value=='Keep note'
    assert w['Finance Input'].sheet_state=='visible'
    assert w['Cash Flow'].sheet_state=='visible'
    assert w['Finance_Data'].sheet_state=='hidden'
    assert not w['Finance Input']['V28'].protection.locked
    assert not package_issues(out)
    w.close()


def test_mapping_preserves_finance(financed_seed,tmp_path):
    from progress_studio.services.mapping_store import MappingStore
    from progress_studio.services.workbook_export_service import WorkbookExportService
    from progress_studio.domain.mapping_models import ActivityRow, BOQRow
    store=MappingStore();store.load_activities([ActivityRow('A1000','1','1.1','Concrete')])
    store.load_boq([BOQRow('BOQ|2','BOQ',2,'1','','','Concrete',1000,'BOQ-ONE')])
    store.selected_activity_ids={'A1000'};store.selected_boq_ids={'BOQ|2'};store.map_selected(100)
    out=tmp_path/'mapped.xlsx';WorkbookExportService().export(financed_seed,out,store)
    src=load_workbook(financed_seed);w=load_workbook(out)
    assert read_finance_inputs(w)==read_finance_inputs(src)
    assert w['Finance Input']['S25'].value=='Keep note'
    assert w['Cash Flow'].sheet_state=='visible'
    assert not package_issues(out)
    w.close();src.close()
