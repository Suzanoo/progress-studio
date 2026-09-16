"""FF-0 engineering checks; expected limitations are NOT acceptance passes."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
import pytest

from tests.fixtures.finance_workbook_probe import (
    OWNED, LAST_ROW, create_seed, extend, snapshot, compare,
    package_issues, opaque_parts, inject_opaque_shape, has_opaque_shape,
    lifecycle_matrix,
)
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook


@pytest.fixture(scope="module")
def seed(tmp_path_factory):
    return create_seed(tmp_path_factory.mktemp("finance-seed"))


def test_extension_preserves_existing_owner_values_formulas_charts_and_policy(seed, tmp_path):
    before = snapshot(seed)
    output = extend(seed, tmp_path / "extended.xlsx")
    after = snapshot(output)
    assert not compare(before, after, before["sheets"])
    assert before["calculation"] == after["calculation"]
    assert all(after["names"].get(k) == v for k, v in before["names"].items())
    assert package_issues(output) == []
    assert len(after["sheets"]["Cash Flow"]["charts"]) == 2
    assert after["sheets"]["Finance Input"]["state"] == "visible"
    assert after["sheets"]["Finance_Data"]["state"] == "hidden"


def test_repeated_refresh_preserves_last_input_row_zero_override_and_custom_notes(seed, tmp_path):
    source = extend(seed, tmp_path / "first.xlsx")
    wb = load_workbook(source)
    inp = wb["Finance Input"]
    inp["B13"] = 0
    inp.cell(LAST_ROW, 1, "LAST-ROW")
    inp.cell(LAST_ROW, 4, 12)
    inp["G25"] = "PO note"
    wb.save(source); wb.close()
    before = snapshot(source)
    for index in range(3):
        target = extend(source, tmp_path / f"refresh-{index}.xlsx", restore_opaque=True)
        after = snapshot(target)
        assert not compare(before, after, before["sheets"])
        assert after["names"] == before["names"]
        assert after["calculation"] == before["calculation"]
        assert package_issues(target) == []
        source = target


def test_capacity_overflow_and_foreign_sheet_collision_fail_without_output(seed, tmp_path):
    source = extend(seed, tmp_path / "source.xlsx")
    wb = load_workbook(source)
    wb["Finance Input"].cell(LAST_ROW + 1, 1, "DO-NOT-DROP")
    wb.save(source); wb.close()
    target = tmp_path / "rejected.xlsx"
    with pytest.raises(ValueError, match="capacity"):
        extend(source, target)
    assert not target.exists()
    wb = load_workbook(source); wb["Finance Input"]["A2"] = "USER OWNED"
    wb.save(source); wb.close()
    with pytest.raises(ValueError, match="refusing replacement"):
        extend(source, target)
    assert not target.exists()


def test_input_cells_unlocked_and_formula_cells_locked_after_refresh(seed, tmp_path):
    output = extend(seed, tmp_path / "inputs.xlsx")
    wb = load_workbook(output)
    try:
        assert wb["Finance Input"].protection.sheet
        for cell in ("B7", "B13", "A20", "E39"):
            assert not wb["Finance Input"][cell].protection.locked
        assert wb["Cash Flow"]["F6"].protection.locked
        assert wb["Cash Flow"].protection.sheet
        assert wb["Finance Input"].tables["FF0_CashEvents"].ref == "A19:E39"
        assert wb["Earned Value"]["M3"].protection.locked is False
        assert wb["Dashboard"]["G5"].protection.locked is False
        assert wb["Dashboard"]["K5"].protection.locked is False
    finally:
        wb.close()


def test_finance_titles_registered_by_ff2_final_policy(seed, tmp_path):
    output = extend(seed, tmp_path / "policy.xlsx")
    wb = load_workbook(output)
    try:
        finalize_workbook(wb)
        # FF-2 registers the public finance titles; FF-0's formulas remain a probe.
        assert wb['Finance Input'].sheet_state == 'visible'
        assert wb['Cash Flow'].sheet_state == 'visible'
        assert wb['Finance_Data'].sheet_state == 'hidden'
    finally:
        wb.close()


def test_opaque_shape_loss_and_existing_restoration_candidate(seed, tmp_path):
    source = inject_opaque_shape(seed, tmp_path / "opaque.xlsx")
    plain = extend(source, tmp_path / "plain.xlsx")
    restored = extend(source, tmp_path / "restored.xlsx", restore_opaque=True)
    assert has_opaque_shape(source)
    assert not has_opaque_shape(plain)  # observed openpyxl limitation
    assert has_opaque_shape(restored)
    assert package_issues(restored) == []
    parts = opaque_parts(restored)
    assert all(parts.get(n) == h for n, h in opaque_parts(source).items())
    assert not compare(snapshot(source), snapshot(restored), snapshot(source)["sheets"])


def test_relationship_validator_rejects_missing_chart_target(seed, tmp_path):
    output = extend(seed, tmp_path / "valid.xlsx")
    bad = tmp_path / "bad.xlsx"
    with ZipFile(output) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    chart = next(n for n in parts if n.startswith("xl/charts/chart") and n.endswith(".xml"))
    del parts[chart]
    with ZipFile(bad, "w", ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    issues = package_issues(bad)
    assert any("missing relationship target" in issue for issue in issues)
    assert any("content type refers to missing part" in issue for issue in issues)


def test_rebuild_matrix_reports_probe_boundaries_after_ff2(seed, tmp_path):
    source = extend(seed, tmp_path / "source.xlsx", restore_opaque=True)
    matrix = lifecycle_matrix(source, tmp_path / "matrix")
    assert set(matrix) == {"progress_snapshot", "progress_live", "payment_snapshot",
                           "payment_live", "earned_value", "payment_breakdown", "mapping_export"}
    assert all(v["status"] != "ERROR" for k,v in matrix.items() if k != "mapping_export"), matrix
    # FF-2 closes visibility registration. The old probe schema is deliberately
    # rejected by Mapping rather than silently converted to production inputs.
    for name in ("progress_snapshot", "progress_live", "payment_snapshot", "payment_live"):
        assert "Finance Input: state changed" not in matrix[name]["issues"], matrix[name]
    assert matrix["earned_value"]["subsequent_finance_refresh"].startswith("PASS")
    assert "charts changed" in matrix["earned_value"]["append_order_negative_probe"]
    assert matrix["mapping_export"]["status"] == "ERROR"
    assert "schema" in str(matrix["mapping_export"])


def test_source_path_and_existing_output_are_never_overwritten(seed, tmp_path):
    original = seed.read_bytes()
    with pytest.raises(ValueError, match="new output path"):
        extend(seed, seed)
    assert seed.read_bytes() == original
    existing = tmp_path / "exists.xlsx"; existing.write_bytes(b"preserve")
    with pytest.raises(ValueError, match="new output path"):
        extend(seed, existing)
    assert existing.read_bytes() == b"preserve"
