from copy import copy
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.workbook.defined_name import DefinedName

from progress_studio.infrastructure.excel.earned_value_workbook import (
    EARNED_VALUE_SHEET, EV_DATA_SHEET, EV_TABLE_SHEET, EV_VIEW_DATE_NAME,
    render_earned_value_sheet,
)
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from tests.unit.earn_value.test_ev_live_dataset import _live_workbook, _result


def _values(ws):
    return {c.coordinate: c.value for row in ws for c in row if c.value is not None}


def _chart_references(ws):
    return [[(s.cat.numRef.f, s.val.numRef.f) for s in chart.series]
            for chart in ws._charts]


@pytest.mark.parametrize("mode", ["snapshot", "live"])
@pytest.mark.parametrize("coordinate", ["M3", "Q7"])
def test_final_policy_unlocks_only_named_ev_control_and_is_repeatable(mode, coordinate):
    wb = Workbook()
    ws = wb.active
    ws.title = EARNED_VALUE_SHEET
    ws["M3"] = datetime(2026, 1, 30)
    ws[coordinate] = datetime(2026, 4, 24)
    ws[coordinate].protection = Protection(locked=True, hidden=True)
    ws["B6"] = "=EV_View_Date"
    wb.defined_names.add(DefinedName(EV_VIEW_DATE_NAME, attr_text=f"'{ws.title}'!{coordinate}"))
    before = _values(ws)
    binding = wb.defined_names[EV_VIEW_DATE_NAME].attr_text
    try:
        for _ in range(2):
            finalize_workbook(wb, mode=mode, include_guide=False)
            assert ws.protection.sheet is True
            assert ws[coordinate].protection.locked is False
            assert ws[coordinate].protection.hidden is True
            assert ws["B6"].protection.locked is True
            assert [c.coordinate for row in ws for c in row if not c.protection.locked] == [coordinate]
            assert _values(ws) == before
            assert wb.defined_names[EV_VIEW_DATE_NAME].attr_text == binding
    finally:
        wb.close()


@pytest.mark.parametrize("reference", [
    None, "#REF!", "42", '"a date"', "OtherName",
    "'Missing Sheet'!$M$3", "'EV_Data'!$M$3", "'[other.xlsx]Earned Value'!$M$3",
    "'Earned Value'!$M$3:$N$4", "'Earned Value'!$M:$M", "'Earned Value'!$3:$3",
    "'Earned Value'!$M$3,'Earned Value'!$N$3", "'Earned Value'!$M$0",
    "'Earned Value'!$XFE$3", "'Earned Value'!$M$1048577", "'Earned Value'!?",
    "OFFSET('Earned Value'!$M$3,0,0)",
    "'Earned Value'!$M$3+1", '"unterminated',
])
def test_missing_or_invalid_ev_binding_never_unlocks_a_fallback_or_range(reference):
    wb = Workbook()
    ws = wb.active
    ws.title = EARNED_VALUE_SHEET
    ws["M3"] = datetime(2026, 4, 24)
    ws["N3"] = "=1+1"
    data = wb.create_sheet(EV_DATA_SHEET)
    data["M3"] = "=EV_View_Date"
    if reference is not None:
        wb.defined_names.add(DefinedName(EV_VIEW_DATE_NAME, attr_text=reference))
    try:
        finalize_workbook(wb, include_guide=False)
        for sheet in wb:
            assert sheet.protection.sheet is True
            assert all(c.protection.locked for row in sheet for c in row)
        if reference is not None:
            assert wb.defined_names[EV_VIEW_DATE_NAME].attr_text == reference
    finally:
        wb.close()


def test_ev_binding_to_a_formula_or_merged_placeholder_stays_locked():
    wb = Workbook()
    ws = wb.active
    ws.title = EARNED_VALUE_SHEET
    ws["M3"] = "=TODAY()"
    ws.merge_cells("Q7:R7")
    try:
        for coordinate in ("M3", "R7"):
            wb.defined_names.add(DefinedName(EV_VIEW_DATE_NAME, attr_text=f"'{ws.title}'!{coordinate}"))
            finalize_workbook(wb, include_guide=False)
            assert ws[coordinate].protection.locked is True
            assert ws["M3"].value == "=TODAY()"
    finally:
        wb.close()


def test_final_policy_without_ev_preserves_existing_controls():
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["G5"] = "Monthly"
    ws["K5"] = datetime(2026, 4, 24)
    ws["M3"] = "=1+1"
    wb.defined_names.add(DefinedName(EV_VIEW_DATE_NAME, attr_text="'Dashboard'!$M$3"))
    try:
        for _ in range(2):
            finalize_workbook(wb, include_guide=False)
            assert ws.protection.sheet
            assert not ws["G5"].protection.locked
            assert not ws["K5"].protection.locked
            assert ws["M3"].protection.locked
            assert EARNED_VALUE_SHEET not in wb.sheetnames
    finally:
        wb.close()


@pytest.mark.parametrize("method", ["rebuild_payment", "rebuild_live_payment"])
def test_payment_final_output_keeps_ev_control_editable_and_ev_content(method, tmp_path):
    source, output = tmp_path / "ev.xlsx", tmp_path / "payment.xlsx"
    wb = _live_workbook()
    for col in range(18, 21):
        wb["main"].cell(3, col, f"W{col - 17}")
    pin = wb.create_sheet("Payment Input")
    for col, value in enumerate(("Type", "WBS", "Activity ID", "Activity Name", "P01"), 1):
        pin.cell(6, col, value)
    pin["E7"] = datetime(2026, 8, 28)
    for col, value in enumerate(("ACT", "1.1.1", "A1", "Activity 1", 0.5), 1):
        pin.cell(8, col, value)
    render_earned_value_sheet(wb, _result())
    wb.save(source)
    wb.close()
    before = load_workbook(source)
    try:
        ev = before[EARNED_VALUE_SHEET]
        assert not ev.protection.sheet and ev["M3"].protection.locked
        expected_values = {name: _values(before[name]) for name in
                           (EARNED_VALUE_SHEET, EV_DATA_SHEET, EV_TABLE_SHEET, "Dashboard_Data", "main")}
        expected_refs = _chart_references(ev)
        expected_validation = copy(ev.data_validations)
        binding = before.defined_names[EV_VIEW_DATE_NAME].attr_text
        formula_cells = [(ws.title, c.coordinate) for ws in before
                         for row in ws for c in row if c.data_type == "f" and c.protection.locked]
    finally:
        before.close()
    getattr(WorkbookRebuildEngine(), method)(source, output)
    after = load_workbook(output)
    try:
        ev = after[EARNED_VALUE_SHEET]
        assert after["Payment"]._images  # real renderer and final save completed
        assert ev.protection.sheet is True
        assert ev["M3"].protection.locked is False
        assert after.defined_names[EV_VIEW_DATE_NAME].attr_text == binding
        assert ev.data_validations == expected_validation
        assert _chart_references(ev) == expected_refs
        for name, expected in expected_values.items():
            assert _values(after[name]) == expected
        assert formula_cells
        assert all(after[name][coord].protection.locked for name, coord in formula_cells)
        # Full calendar extends beyond the selected view / Actual Data Cutoff.
        assert after["Dashboard_Data"]["K4"].value > ev["M3"].value
        assert ev.data_validations.dataValidation[0].formula1.endswith('$K$2:$K$4")')
    finally:
        after.close()
