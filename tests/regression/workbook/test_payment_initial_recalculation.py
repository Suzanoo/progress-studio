"""Package contract only: Microsoft Excel first-open rendering needs manual acceptance."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from progress_studio.services.payment_service import PaymentService
from tests.integration.mapping.test_mapping_overlay_export import _created_mapping_input

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def _seed_calculated_formula(path):
    wb = load_workbook(path)
    wb.create_sheet('Cache Probe')['A1'] = '=1+1'
    wb.save(path)
    wb.close()
    # Model an Excel-calculated input, not a cache restoration in production.
    with ZipFile(path) as z:
        parts = [(i, z.read(i.filename)) for i in z.infolist()]
    seeded = False
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for info, raw in parts:
            if info.filename.startswith('xl/worksheets/sheet') and info.filename.endswith('.xml'):
                root = ET.fromstring(raw)
                for cell in root.findall('.//s:c', NS):
                    if cell.findtext('s:f', namespaces=NS) == '1+1':
                        cell.find('s:v', NS).text = '2'
                        raw = ET.tostring(root)
                        seeded = True
            z.writestr(info, raw)
    assert seeded


def _content(wb):
    return {ws.title: {c.coordinate: c.value for row in ws for c in row if c.value is not None}
            for ws in wb if ws.title not in ('README', 'Payment')}


def _charts(wb):
    return {ws.title: [[(s.cat.numRef.f, s.val.numRef.f) for s in chart.series]
                       for chart in ws._charts] for ws in wb if ws._charts}


@pytest.mark.parametrize('method', ['rebuild_payment', 'rebuild_live_payment'])
def test_payment_output_requests_manual_initial_calculation_after_each_final_save(tmp_path, method):
    source, _ = _created_mapping_input(tmp_path, True)
    _seed_calculated_formula(source)
    cached = load_workbook(source, data_only=True)
    assert cached['Cache Probe']['A1'].value == 2
    cached.close()
    wb = load_workbook(source)
    expected, charts = _content(wb), _charts(wb)
    wb.close()
    for iteration in range(2):
        output = tmp_path / f'payment-{iteration}.xlsx'
        getattr(WorkbookRebuildEngine(), method)(source, output)
        wb = load_workbook(output)
        try:
            assert _content(wb) == expected
            assert _charts(wb) == charts
            assert wb['Payment']._images
            calc = wb.calculation
            assert calc.calcMode == 'manual' and calc.calcOnSave
            assert calc.fullCalcOnLoad and calc.forceFullCalc and calc.calcId == 0
        finally:
            wb.close()
        with ZipFile(output) as z:
            calc = ET.fromstring(z.read('xl/workbook.xml')).find('s:calcPr', NS)
            assert calc.get('calcMode') == 'manual'
            assert calc.get('fullCalcOnLoad') == calc.get('forceFullCalc') == '1'
            assert calc.get('calcId') == '0'
        cached = load_workbook(output, data_only=True)
        # Python still does not calculate: do not mistake the request for visual acceptance.
        assert cached['Cache Probe']['A1'].value is None
        cached.close()
        source = output


def test_standalone_payment_keeps_existing_save_policy(tmp_path):
    source, _ = _created_mapping_input(tmp_path, True)
    output = tmp_path / 'standalone.xlsx'
    PaymentService().render_payment_backbones(source, source, output)
    wb = load_workbook(output)
    try:
        assert wb.calculation.calcMode == 'manual'
        assert wb.calculation.fullCalcOnLoad is False
        assert wb.calculation.forceFullCalc is False
    finally:
        wb.close()
