from datetime import date
import pytest
from openpyxl import load_workbook

from progress_studio.app.desktop import DesktopRunner, DesktopRunOptions
from progress_studio.domain.mapping_models import BOQRow
from progress_studio.infrastructure.schedule_xml import NormalizedScheduleXmlReader
from progress_studio.infrastructure.excel.weight_basis import creation_basis,creation_field,uses_dummy_weights
from progress_studio.infrastructure.excel.mapping_reader import ProgressActivityReader
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.earned_value_input_reader import EarnedValueInputWorkbookReader,EarnedValueWorkbookInputError
from progress_studio.infrastructure.excel.finance_input_workbook import read_finance_inputs
from progress_studio.infrastructure.excel.payment_breakdown_workbook import render_payment_breakdown
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.services.mapping_store import MappingStore
from progress_studio.services.workbook_export_service import WorkbookExportService
from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from progress_studio.services.earned_value_rebuild_service import EarnedValueRebuildService
from progress_studio.services.financial_forecast_workbook_service import FinancialForecastWorkbookService
from progress_studio.services.financial_forecast_deriver import FinancialForecastDeriver
from progress_studio.services.payment_breakdown_adapter import MainDatasetPaymentBreakdownAdapter
from tests.fixtures.amount_xml import amount_xml


def test_amount_create_finance_mapping_ev_and_rebuild_share_existing_contract(tmp_path,monkeypatch):
    xml=amount_xml(tmp_path,'p6')
    field=NormalizedScheduleXmlReader().read_with_amount_fields(xml)[2][0]
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path',lambda:tmp_path/'desktop')
    source=DesktopRunner().run(DesktopRunOptions(xml,'5',weight_basis='amount',amount_field=field.identity)).output_workbook
    with pytest.raises(EarnedValueWorkbookInputError):
        EarnedValueInputWorkbookReader().read(source)
    wb=load_workbook(source)
    snapshot=MainDatasetPaymentBreakdownAdapter().derive(main_dataset_from_workbook(wb),min_occurrences=1)
    assert snapshot.eligible_source_count==2
    assert set(snapshot.skipped_activity_ids)=={'A3','A4'}
    render_payment_breakdown(wb,snapshot);finalize_workbook(wb)
    assert 'dummy' not in wb['Payment-Breakdown']['A1'].value.lower()
    assert 'monetary XML values' in wb['Payment-Breakdown']['D4'].comment.text
    wb.save(source);wb.close()
    finance=tmp_path/'finance.xlsx'
    FinancialForecastWorkbookService().generate(source,finance,opening_date=date(2026,1,1),actuals_through=date(2026,1,31),currency='THB',capacity_rows=2)
    wb=load_workbook(finance)
    expected_finance=read_finance_inputs(wb)
    derived=FinancialForecastDeriver().derive(expected_finance)
    assert not derived.items and derived.remaining_receivable==0 and derived.peak_funding_requirement==0
    assert creation_field(wb)==field
    wb.close()
    store=MappingStore();store.load_activities(ProgressActivityReader().read(finance))
    store.load_boq([BOQRow('B1','BOQ',2,'1','','','Contract item',1000,'BOQ-1')])
    store.selected_activity_ids={'A1'};store.selected_boq_ids={'B1'};store.map_selected(100)
    mapped=tmp_path/'mapped.xlsx'
    result=WorkbookExportService().export(finance,mapped,store)
    assert result.validation.is_complete
    assert EarnedValueInputWorkbookReader().read(mapped).boq_rows[0].amount==1000
    ev=tmp_path/'ev.xlsx';EarnedValueRebuildService().generate(mapped,ev)
    paths=[mapped,ev]
    for method in ('rebuild_progress','rebuild_live_progress'):
        output=tmp_path/f'{method}.xlsx'
        getattr(WorkbookRebuildEngine(),method)(ev,output)
        paths.append(output)
    for path in paths:
        wb=load_workbook(path)
        try:
            assert creation_basis(wb)=='amount' and creation_field(wb)==field
            assert not uses_dummy_weights(wb)
            assert 'Current amounts: BOQ Mapping' in wb['README']['B4'].value
            assert sum(a.amount for a in main_dataset_from_workbook(wb).activities)==1000
            assert read_finance_inputs(wb)==expected_finance
            assert wb['main'].protection.sheet and wb['Dashboard'].protection.sheet
            if path!=mapped:assert 'Earned Value' in wb.sheetnames
        finally:wb.close()
