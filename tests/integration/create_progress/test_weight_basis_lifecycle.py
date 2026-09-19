import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from progress_studio.app.desktop import DesktopRunner, DesktopRunOptions
from progress_studio.infrastructure.excel.import_workbook_writer import ImportWorkbookWriter
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.weight_basis import creation_basis, set_creation_basis
from progress_studio.infrastructure.schedule_xml import NormalizedScheduleXmlReader
from progress_studio.services.import_service import ImportService
from progress_studio.services.schedule_service import ScheduleService
from progress_studio.services.distribution import get_distribution
from progress_studio.services.rebuild_service import WorkbookRebuildEngine
from progress_studio.services.workbook_export_service import WorkbookExportService
from tests._paths import FIXTURES_ROOT
from tests.integration.mapping.test_mapping_overlay_export import _created_mapping_input


def weighted_xml(tmp_path, source, bad_duration=False):
    # Adapt existing source fixtures; dates deliberately do not imply these durations.
    filename = 'p6_n5.xml' if source == 'p6' else 'msp_n3.xml'
    tree = ET.parse(FIXTURES_ROOT / 'xml' / filename)
    root = tree.getroot()
    ns = root.tag.split('}')[0] + '}'
    if source == 'p6':
        parent = root.find(ns + 'Project')
        activities = parent.findall(ns + 'Activity')
        for activity, hours in zip(activities, ('8.5', '40')):
            ET.SubElement(activity, ns + 'PlannedDuration').text = hours
            ET.SubElement(activity, ns + 'Type').text = 'Task Dependent'
        milestone = ET.SubElement(parent, ns + 'Activity')
        fields = dict(Id='M1', Name='Handover', WBSObjectId='30337',
                      PlannedStartDate='2026-11-16T17:00:00', PlannedFinishDate='2026-11-16T17:00:00',
                      PlannedDuration='0', Type='Finish Milestone')
        duration_tag = 'PlannedDuration'
    else:
        parent = root.find(ns + 'Tasks')
        activities = [t for t in parent if t.findtext(ns + 'Summary') == '0']
        for activity, duration in zip(activities, ('PT8H30M', 'PT40H')):
            ET.SubElement(activity, ns + 'Duration').text = duration
            ET.SubElement(activity, ns + 'Cost').text = '999999'
        milestone = ET.SubElement(parent, ns + 'Task')
        fields = dict(UID='7', ID='7', Name='Handover', WBS='2.1.1.1', OutlineLevel='5', Summary='0',
                      Start='2026-11-16T17:00:00', Finish='2026-11-16T17:00:00', Duration='PT0H', Milestone='1')
        duration_tag = 'Duration'
    for key, value in fields.items():
        ET.SubElement(milestone, ns + key).text = value
    if bad_duration:
        activities[0].find(ns + duration_tag).text = 'invalid'
    path = tmp_path / (source + '.xml')
    tree.write(path, encoding='utf-8', xml_declaration=True)
    return path


@pytest.mark.parametrize('source', ['p6', 'msp'])
@pytest.mark.parametrize('basis', ['equal', 'duration'])
def test_source_weights_and_invalid_duration_stop_before_output(tmp_path, source, basis):
    reader = NormalizedScheduleXmlReader()
    service = ImportService(reader, ScheduleService(), ImportWorkbookWriter())
    source_xml = weighted_xml(tmp_path, source)
    _, rows = reader.read(source_xml)
    activities = [r for r in rows if not r.is_summary]
    assert [r.duration_hours for r in activities] == [8.5, 40, 0]
    assert [r.is_milestone for r in activities] == [False, False, True]
    assert all(r.amount is None for r in activities)
    output = tmp_path / 'imported.xlsx'
    service.import_xml(source_xml, output, weight_basis=basis)
    wb = load_workbook(output)
    try:
        assert creation_basis(wb) == basis
        ws = wb['main']; headers = {c.value: c.column for c in ws[1]}
        weights = [ws.cell(r, headers['XML Amount']).value for r in range(2, ws.max_row + 1)
                   if ws.cell(r, headers['Row Type']).value == 'Activity']
        assert weights == ([1, 1, 0] if basis == 'equal' else [8.5, 40, 0])
    finally:
        wb.close()
    bad_xml = weighted_xml(tmp_path, source, bad_duration=True)
    invalid_output = tmp_path / 'invalid.xlsx'
    with pytest.raises(ValueError, match='source duration'):
        service.import_xml(bad_xml, invalid_output, weight_basis='duration')
    assert not invalid_output.exists()


@pytest.mark.parametrize('method', ['auto', 'flat', 'front', 'back', 'bell'])
def test_weighting_preserves_plan_distribution_and_default(tmp_path, monkeypatch, method):
    source = weighted_xml(tmp_path, 'msp')
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path', lambda: tmp_path / 'output')
    # Default Equal is exercised through the complete Desktop Create pipeline.
    result = DesktopRunner().run(DesktopRunOptions(source, '5', distribution_method=method))
    wb = load_workbook(result.output_workbook)
    try:
        dataset = main_dataset_from_workbook(wb)
        assert creation_basis(wb) == 'equal'
        assert [a.amount for a in dataset.activities] == [1, 1, 0]
        assert dataset.activities[-1].description == 'Handover'
        for activity in dataset.activities[:2]:
            assert sum(value or 0 for _, value in activity.period_values) == pytest.approx(1)
        assert wb['main'].protection.sheet
        assert wb.calculation.calcMode == 'manual'
        assert len(wb['main']._charts) == 1
        assert len(wb['Dashboard']._charts) == 1
        # Distribution request remains independent and recorded by the existing engine.
        expected_mode = 'AUTO' if method == 'auto' else get_distribution(method).name
        assert any(expected_mode == c.value for row in wb['Info'] for c in row)
    finally:
        wb.close()


@pytest.mark.parametrize('basis', ['equal', 'duration'])
@pytest.mark.parametrize('live', [False, True])
def test_mapping_regeneration_and_progress_rebuild_preserve_basis(tmp_path, basis, live):
    source, store = _created_mapping_input(tmp_path, populated_payment=True)
    wb = load_workbook(source)
    set_creation_basis(wb, basis)
    wb.save(source); wb.close()
    mapped = tmp_path / 'mapped.xlsx'
    WorkbookExportService().export(source, mapped, store)
    wb = load_workbook(mapped)
    try:
        assert creation_basis(wb) == basis
        assert 'Current amounts: BOQ Mapping' in wb['README']['B4'].value
        assert main_dataset_from_workbook(wb).activities[0].amount == 750
        assert wb['Payment']._images
    finally:
        wb.close()
    rebuilt = tmp_path / 'rebuilt.xlsx'
    engine = WorkbookRebuildEngine()
    rebuild = engine.rebuild_live_progress if live else engine.rebuild_progress
    rebuild(mapped, rebuilt)
    wb = load_workbook(rebuilt)
    try:
        assert creation_basis(wb) == basis
        assert 'Current amounts: BOQ Mapping' in wb['README']['B4'].value
        assert main_dataset_from_workbook(wb).activities[0].amount == 750
        assert wb['Payment']._images
        assert wb['main'].protection.sheet
        assert len(wb['Dashboard']._charts) == 1
    finally:
        wb.close()


@pytest.mark.parametrize('basis', ['equal', 'duration'])
@pytest.mark.parametrize('live', [False, True])
def test_dummy_rebuild_keeps_weights_zero_milestone_and_labels(tmp_path, monkeypatch, basis, live):
    source = weighted_xml(tmp_path, 'msp')
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path', lambda: tmp_path / 'output')
    result = DesktopRunner().run(DesktopRunOptions(source, '5', weight_basis=basis))
    output = tmp_path / 'rebuilt.xlsx'
    engine = WorkbookRebuildEngine()
    rebuild = engine.rebuild_live_progress if live else engine.rebuild_progress
    rebuild(result.output_workbook, output)
    wb = load_workbook(output)
    try:
        assert creation_basis(wb) == basis
        assert 'dummy units, not Contract Value' in wb['README']['B4'].value
        assert [a.amount for a in main_dataset_from_workbook(wb).activities] == (
            [1, 1, 0] if basis == 'equal' else [8.5, 40, 0])
        if live:
            assert wb['Dashboard']['J6'].value == 'Total weight'
    finally:
        wb.close()


def test_dummy_is_not_ev_budget_or_finance_cash(tmp_path):
    from datetime import date
    from progress_studio.infrastructure.excel.earned_value_input_reader import (
        EarnedValueInputWorkbookReader, EarnedValueWorkbookInputError,
    )
    from progress_studio.infrastructure.excel.finance_input_workbook import read_finance_inputs
    from progress_studio.services.financial_forecast_workbook_service import FinancialForecastWorkbookService
    from progress_studio.services.financial_forecast_deriver import FinancialForecastDeriver
    from tests.integration.rebuild.test_progress_rebuild_engine import _full_rebuild_fixture

    source = _full_rebuild_fixture(tmp_path / 'dummy.xlsx')
    wb = load_workbook(source)
    set_creation_basis(wb, 'equal')
    wb.save(source); wb.close()
    with pytest.raises(EarnedValueWorkbookInputError):
        EarnedValueInputWorkbookReader().read(source)
    output = tmp_path / 'finance.xlsx'
    FinancialForecastWorkbookService().generate(
        source, output, opening_date=date(2026, 1, 1),
        actuals_through=date(2026, 1, 31), currency='THB', capacity_rows=2,
    )
    wb = load_workbook(output)
    try:
        assert creation_basis(wb) == 'equal'
        finance = FinancialForecastDeriver().derive(read_finance_inputs(wb))
        assert finance.remaining_receivable == 0
        assert finance.peak_funding_requirement == 0
        assert not finance.items
        assert wb['main']['F11'].value == 1000  # progress never used as cash
    finally:
        wb.close()
