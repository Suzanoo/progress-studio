from datetime import datetime
from dataclasses import replace
from pathlib import Path
import pytest
from openpyxl import load_workbook
from tests.unit.earn_value.test_ev_live_dataset import _live_workbook
from progress_studio.infrastructure.excel.ev_monetary_inputs import ACTIVITY, BOQ, SHEET, write_inputs, monetary
from progress_studio.services.earned_value_rebuild_service import EarnedValueRebuildService, EarnedValueRebuildError
from progress_studio.infrastructure.excel.rebuild_workbook_reader import RebuildWorkbookReader


def source_book(path):
    wb = _live_workbook()
    main = wb['main']
    for col, key in enumerate(('W1','W2','W3'),18):
        main.cell(3,col,key)
    main['O9'] = 1
    main['R9'] = 1; main['S9'] = 0; main['T9'] = 0
    main['R10'] = 1; main['S10'] = 0; main['T10'] = 0
    for row, pa, aid in ((11,'P','A2'),(12,'A','A2'),(13,'P','M1'),(14,'A','M1')):
        main.cell(row,1,'Activity' if pa=='P' else None)
        main.cell(row,2,'1.1.2' if aid=='A2' else '1.1.3')
        main.cell(row,3,aid);main.cell(row,4,pa);main.cell(row,5,aid)
        main.cell(row,15,1 if aid=='A2' and pa=='P' else 0)
        for col in range(18,21):main.cell(row,col,0)
    main['T11'] = 1
    del wb['BOQ Activity Mapping']
    write_inputs(wb,[('A1','First',100,False,None,None),('A2','Second',900,False,None,None),
                     ('M1','Milestone',200,True,datetime(2026,4,1),datetime(2026,8,1))])
    wb.save(path);wb.close()
    return path


def add_boq(path, mismatch=False):
    wb=load_workbook(path)
    ws=wb.create_sheet('BOQ Activity Mapping')
    ws.append(['Activity ID','BOQ Key','Source Sheet','Source Row','WBS-2','WBS-3','WBS-4','BOQ Description','BOQ Amount','Share %','Allocated Amount','Mapping ID','BOQ ID'])
    for i,(aid,value) in enumerate((('A1',100),('A2',900),('M1',200)),1):
        ws.append([aid,f'B{i}','Contract',i,'Work','','',aid,value,1,value+1 if mismatch and i==1 else value,f'M{i}',f'BOQ{i}'])
    ws=wb.create_sheet('Mapping Summary')
    for record in [('BOQ items',3),('Fully allocated BOQ items',3),('Partially allocated BOQ items',0),('Unmapped BOQ items',0),('Allocated percent',1)]:ws.append(record)
    wb['main']['O9']=100;wb['main']['O11']=900
    wb.save(path);wb.close()


def test_activity_source_aggregates_money_and_preserves_live_formulas(tmp_path):
    source=source_book(tmp_path/'source.xlsx'); output=tmp_path/'ev.xlsx'
    service=EarnedValueRebuildService()
    result=service.generate(source,output,monetary_source=ACTIVITY)
    assert result.project_bac==1200
    derived,*_=service._derive(source, ACTIVITY)
    assert [p.earned_value for p in derived.project_points]==[100,100,300]
    assert [p.planned_value for p in derived.project_points]==[100,300,1200]
    wb=load_workbook(output)
    assert 'BOQ Activity Mapping' not in wb.sheetnames
    assert wb['EV Table']['A5'].value=='Activity ID'
    assert wb['EV Table']['H3'].value==ACTIVITY
    assert wb['EV_Data']['BE2'].value==100
    assert wb['EV_Data']['BE3'].value==900
    assert wb['EV_Data']['BG4'].value=="=IF(ISNUMBER('EV Monetary Inputs'!$F$6),IF(INT('EV Monetary Inputs'!$F$6)<=EV_View_Date,1,0),0)"
    assert wb['EV_Data']['BQ9'].value==100 and wb['EV_Data']['BR10'].value==100
    assert wb['EV_Data']['BQ13'].value==0
    assert wb['Earned Value']['G6'].value=='=SUM(EV_Data!BI2:BI4)'
    assert wb['main']['O13'].value==0
    assert wb['EV_Data']['AG2'].value=='=BH2'
    assert 'SUMPRODUCT(main!' in wb['EV_Data']['BY2'].value
    wb['EV Monetary Inputs']['C4']=150
    wb['Earned Value']['M3']=datetime(2026,4,24)
    wb.save(output);wb.close()
    refreshed=tmp_path/'refreshed.xlsx'
    assert service.generate(output,refreshed,monetary_source=ACTIVITY).project_bac==1250
    wb=load_workbook(refreshed)
    assert wb['Earned Value']['M3'].value==datetime(2026,4,24)
    assert wb['EV_Data']['BE2'].value==150
    wb.close()


@pytest.mark.parametrize('bad',[None,-1,'text','NaN','Infinity','1e309','1e-320',True,'=1+2'])
def test_invalid_amount_blocks(tmp_path,bad):
    path=source_book(tmp_path/'input.xlsx');wb=load_workbook(path)
    wb[SHEET]['C4']=bad;wb.save(path);wb.close()
    with pytest.raises(EarnedValueRebuildError):EarnedValueRebuildService().analyze(path,monetary_source=ACTIVITY)


def test_zero_duplicate_summary_and_independent_boq(tmp_path):
    path=source_book(tmp_path/'input.xlsx');wb=load_workbook(path)
    wb[SHEET]['C4']=0
    wb['main']['O5']=99999999
    wb.create_sheet('BOQ Activity Mapping')['A1']='broken'
    wb.create_sheet('Mapping Summary')['A1']='incomplete'
    wb.save(path);wb.close()
    assert EarnedValueRebuildService().analyze(path,monetary_source=ACTIVITY).project_bac==1100
    wb=load_workbook(path)
    wb[SHEET]['C5']=0;wb[SHEET]['C6']=0;wb.save(path);wb.close()
    with pytest.raises(EarnedValueRebuildError,match='positive'):EarnedValueRebuildService().analyze(path,monetary_source=ACTIVITY)
    wb=load_workbook(path);wb['main']['E11']='A1';wb.save(path);wb.close()
    with pytest.raises(EarnedValueRebuildError,match='Duplicate'):EarnedValueRebuildService().analyze(path,monetary_source=ACTIVITY)


def test_boq_reconciliation_and_source_switch(tmp_path):
    path=source_book(tmp_path/'input.xlsx');add_boq(path)
    service=EarnedValueRebuildService();a=tmp_path/'a.xlsx';b=tmp_path/'b.xlsx';c=tmp_path/'c.xlsx'
    service.generate(path,a,monetary_source=ACTIVITY)
    service.generate(a,b,monetary_source=BOQ)
    wb=load_workbook(b)
    assert wb['EV Table']['A5'].value=='BOQ ID'
    assert wb['EV_Data']['BL2'].value==100
    wb['EV Monetary Inputs']['C4']=500;wb.save(b);wb.close()
    assert service.analyze(b,monetary_source=ACTIVITY).project_bac==1600
    assert service.analyze(b,monetary_source=BOQ).project_bac==1200
    service.generate(b,c,monetary_source=ACTIVITY)
    wb=load_workbook(c)
    assert wb['EV_Data']['BL2'].value is None
    assert wb['EV Table']['A5'].value=='Activity ID';wb.close()
    wb=load_workbook(path);wb['main']['O9']=99;wb.save(path);wb.close()
    with pytest.raises(EarnedValueRebuildError,match='reconciliation mismatch'):service.analyze(path,monetary_source=BOQ)
    other=source_book(tmp_path/'mismatch.xlsx');add_boq(other,True)
    with pytest.raises(EarnedValueRebuildError,match='reconciliation mismatch'):service.analyze(other,monetary_source=BOQ)


def test_milestone_date_boundaries_and_weight(tmp_path):
    path=source_book(tmp_path/'input.xlsx');service=EarnedValueRebuildService()
    wb=load_workbook(path);wb[SHEET]['E6']=datetime(2026,4,24);wb[SHEET]['F6']=datetime(2026,4,24);wb.save(path);wb.close()
    result,*_=service._derive(path,ACTIVITY)
    milestone=result.activities[-1]
    assert [p.planned_value for p in milestone.points]==[0,200,200]
    assert [p.earned_value for p in milestone.points]==[0,200,200]
    wb=load_workbook(path);wb[SHEET]['F6']=None;wb.save(path);wb.close()
    result,*_=service._derive(path,ACTIVITY)
    assert [p.earned_value for p in result.activities[-1].points]==[0,0,0]
    wb=load_workbook(path);wb['main']['O13']=200;wb.save(path);wb.close()
    with pytest.raises(EarnedValueRebuildError,match='Progress Weight'):service.analyze(path,monetary_source=ACTIVITY)


@pytest.mark.parametrize('source',['p6','msp'])
@pytest.mark.parametrize('basis',['equal','duration','amount'])
def test_xml_create_bac_and_progress_rebuild_preservation(tmp_path,monkeypatch,source,basis):
    from tests.fixtures.amount_xml import amount_xml
    from progress_studio.app.desktop import DesktopRunner,DesktopRunOptions
    from progress_studio.infrastructure.schedule_xml import NormalizedScheduleXmlReader
    from progress_studio.services.rebuild_service import WorkbookRebuildEngine
    xml=amount_xml(tmp_path,source,('100','900','0','200'))
    field=NormalizedScheduleXmlReader().read_with_amount_fields(xml)[2][0]
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path',lambda:tmp_path/'desktop')
    created=DesktopRunner().run(DesktopRunOptions(xml,'5',weight_basis=basis,amount_field=field.identity if basis=='amount' else None,distribution_method='flat')).output_workbook
    wb=load_workbook(created)
    assert wb[SHEET]['D7'].value is True
    assert wb[SHEET]['C7'].value == (200 if basis=='amount' else None)
    if basis!='amount':
        for row,value in enumerate((100,900,0,200),4):wb[SHEET].cell(row,3,value)
        wb.save(created)
    wb.close()
    service=EarnedValueRebuildService()
    assert service.analyze(created,monetary_source=ACTIVITY).project_bac==1200
    ev=tmp_path/'ev.xlsx';service.generate(created,ev,monetary_source=ACTIVITY)
    for live in (False,True):
        out=tmp_path/f'rebuild-{live}.xlsx';engine=WorkbookRebuildEngine()
        method=engine.rebuild_live_progress if live else engine.rebuild_progress
        method(ev,out)
        wb=load_workbook(out)
        assert wb[SHEET]['C7'].value==200
        assert wb[SHEET]['F7'].protection.locked is False
        assert wb['Earned Value']['G6'].value=='=SUM(EV_Data!BI2:BI5)'
        wb.close()
        assert service.analyze(out,monetary_source=ACTIVITY).project_bac==1200


def test_mapping_regeneration_preserves_independent_bac_and_zero_milestone(tmp_path,monkeypatch):
    from tests.fixtures.amount_xml import amount_xml
    from progress_studio.app.desktop import DesktopRunner,DesktopRunOptions
    from progress_studio.infrastructure.schedule_xml import NormalizedScheduleXmlReader
    from progress_studio.domain.mapping_models import BOQRow
    from progress_studio.infrastructure.excel.mapping_reader import ProgressActivityReader
    from progress_studio.services.mapping_store import MappingStore
    from progress_studio.services.workbook_export_service import WorkbookExportService
    from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
    xml=amount_xml(tmp_path,'p6',('100','900','0','200'))
    field=NormalizedScheduleXmlReader().read_with_amount_fields(xml)[2][0]
    monkeypatch.setattr('progress_studio.pipeline.import_step.desktop_path',lambda:tmp_path/'desktop')
    source=DesktopRunner().run(DesktopRunOptions(xml,'5',weight_basis='amount',amount_field=field.identity)).output_workbook
    store=MappingStore();store.load_activities(ProgressActivityReader().read(source))
    store.load_boq([BOQRow('B1','BOQ',2,'1','','','Ordinary',1000,'BOQ1'),BOQRow('B2','BOQ',3,'1','','','Milestone',500,'BOQ2')])
    for aid,bid in [('A1','B1'),('A4','B2')]:
        store.selected_activity_ids={aid};store.selected_boq_ids={bid};store.map_selected(100)
    mapped=tmp_path/'mapped.xlsx';WorkbookExportService().export(source,mapped,store)
    wb=load_workbook(mapped)
    assert [a.amount for a in main_dataset_from_workbook(wb).activities]==[1000,0,0,0]
    assert [wb[SHEET].cell(r,3).value for r in range(4,8)]==[100,900,0,200]
    wb.close()
    service=EarnedValueRebuildService()
    assert service.analyze(mapped,monetary_source=ACTIVITY).project_bac==1200
    assert service.analyze(mapped,monetary_source=BOQ).project_bac==1500
