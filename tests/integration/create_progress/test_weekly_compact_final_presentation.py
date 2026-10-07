from copy import copy
from datetime import datetime
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter as L

from scripts.generate_weekly_compact_final_acceptance import build_acceptance
from progress_studio.infrastructure.excel.live_monthly_workbook import _display_month_buckets
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.weekly_compact_workbook import build_weekly_compact
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook


@pytest.fixture(scope="module")
def artifact(tmp_path_factory):
    folder=tmp_path_factory.mktemp("compact-final")
    path=build_acceptance(folder/"acceptance.xlsx")
    return path


def test_create_projects_monthly_body_footer_and_all_row_scopes(artifact):
    w=load_workbook(artifact);s=w['Weekly Compact'];m=w['main_monthly']
    buckets=_display_month_buckets(w['main'])
    source={(c.value.year,c.value.month):c.column for c in m[4] if isinstance(c.value,datetime)}
    # Project, Activity, WBS, nested WBS and all P/A/footer rows, including zeros
    # and absent observations, share the same direct monthly output contract.
    for key,cols in buckets:
        for r in list(range(5,35))+list(range(36,40)):
            ref=f"'main_monthly'!{L(source[key])}{r}"
            cell=s.cell(r,cols[0])
            assert cell.value==f'=IF({ref}="","",{ref})'
            assert cell.number_format==m.cell(r,source[key]).number_format
            assert copy(cell.font)==copy(m.cell(r,source[key]).font)
            assert not cell.alignment.shrinkToFit and cell.protection.locked
            assert all(isinstance(s.cell(r,c),MergedCell) for c in cols[1:])
    assert w['main']['FG18'].value==0  # Actual zero input, not missing data.
    assert w['main']['FG30'].value is None  # No Actual observation.
    assert 'COUNT(' in m['AZ18'].value and 'COUNT(' in m['AZ30'].value
    assert s['FG39'].value=='=IF(\'main_monthly\'!AZ39="","",\'main_monthly\'!AZ39)'
    # The normal Create renderer owns its own accepted last-nonblank footer.
    assert 'LOOKUP' in m['AZ39'].value
    assert 'PS_WEEKLY_OVERLAY_CUTOFF' not in s['FG39'].value


def test_merged_band_cf_anchors_and_geometry(artifact):
    w=load_workbook(artifact);s=w['Weekly Compact']
    oct_rules=[rules for area,rules in s.conditional_formatting._cf_rules.items() if str(area.sqref)=='FG5:FK34'][0]
    assert len(oct_rules)==8
    assert all('$FG5<>""' in r.formula[0] for r in oct_rules)
    assert 'FG13:FK13' in s.merged_cells and 'FG14:FK14' in s.merged_cells
    assert 'X36:AA36' in s.merged_cells # leap February, four weeks
    assert 'AB36:AF36' in s.merged_cells # March, five weeks
    assert 'T36:W36' in s.merged_cells # first partial month
    assert 'FY36:GB36' in s.merged_cells # final partial month
    assert 'BP36:BS36' in s.merged_cells and 'BT36:BX36' in s.merged_cells
    assert all(s.column_dimensions[L(c)].width==2.5 for c in range(18,186))


def test_checkpoints_use_weekly_values_and_existing_selection(artifact):
    w=load_workbook(artifact);data=w['Dashboard_Data'];chart=w['Weekly Compact']._charts[0]
    assert len(chart.series)==7
    assert chart.series[0].val.numRef.f=="'Dashboard_Data'!$U$2:$U$164"
    assert chart.series[1].val.numRef.f=="'Dashboard_Data'!$V$2:$V$164"
    assert all(s.cat==chart.series[0].cat for s in chart.series)
    assert data['T148'].value==datetime(2026,10,23)
    assert data['T149'].value==datetime(2026,10,30)
    assert 'COUNTIF($K$2:' in data['AC149'].value
    assert 'NOT(IFERROR(W149=1,FALSE))' in data['AD149'].value
    assert 'COUNT($C$2:C148)>0' in data['AE149'].value
    assert 'V149' in data['AE149'].value
    assert 'IFERROR(W148=1,FALSE)' in data['AF148'].value
    assert 'U148' in data['AF148'].value and 'V148' in data['AG148'].value
    assert all(data.cell(2,c).value=='=NA()' for c in range(29,34))
    assert [s.marker.size for s in chart.series[3:]]==[4,4,8,8]
    for s in chart.series[-2:]:
        assert s.marker.graphicalProperties.solidFill.srgbClr=='7030A0'
        assert s.dLbls.spPr.solidFill.srgbClr=='F2EAF8'
        assert s.graphicalProperties.line.noFill
    assert chart.series[2].errBars==w['main']._charts[0].series[2].errBars


def test_dynamic_readonly_date_and_main_input(artifact):
    w=load_workbook(artifact);s=w['Weekly Compact']
    assert next(w.defined_names['PS_WEEKLY_OVERLAY_CUTOFF'].destinations)==('main','$M$35')
    assert not w['main']['M35'].protection.locked
    assert s['C40'].value=='=PS_WEEKLY_OVERLAY_CUTOFF' and s['C40'].protection.locked
    assert s['M35'].value is None and not s.data_validations.dataValidation
    # Shift the finalized layout, then regenerate from new semantic row identities.
    for name in ['main','main_monthly']:w[name].insert_rows(5,4)
    build_weekly_compact(w,main_dataset_from_workbook(w))
    assert w['Weekly Compact']['C44'].value=='=PS_WEEKLY_OVERLAY_CUTOFF'


def test_preserving_finalizer_retains_complete_chart_and_merges(artifact,tmp_path):
    w=load_workbook(artifact);before=set(map(str,w['Weekly Compact'].merged_cells))
    finalize_workbook(w,mode='live');out=tmp_path/'saved.xlsx';w.save(out)
    z=load_workbook(out)
    assert set(map(str,z['Weekly Compact'].merged_cells))==before
    assert len(z['Weekly Compact']._charts[0].series)==7
    assert z['Weekly Compact']['C40'].protection.locked
    ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with ZipFile(out) as package:
        charts=[ET.fromstring(package.read(n)) for n in package.namelist() if n.startswith('xl/charts/chart') and n.endswith('.xml')]
        compact=next(c for c in charts if len(c.findall('.//c:ser',ns))==7)
        assert compact.find('.//c:plotArea/c:spPr/a:noFill',ns) is not None
        assert compact.find('./c:spPr/a:noFill',ns) is not None
