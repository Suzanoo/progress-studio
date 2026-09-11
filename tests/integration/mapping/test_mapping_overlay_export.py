from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple

from progress_studio.domain.mapping_models import ActivityRow, BOQRow
from progress_studio.infrastructure.excel.payment_input_workbook import PaymentInputWorkbook
from progress_studio.services.mapping_store import MappingStore
from progress_studio.services.workbook_export_service import WorkbookExportService
from progress_studio.services.workbook_generation_service import WorkbookGenerationService
from progress_studio.services.working_tree_schedule_source import WorkingTreeScheduleSource
from tests.integration.rebuild.test_progress_rebuild_engine import _full_rebuild_fixture


def _created_mapping_input(tmp_path: Path, populated_payment: bool):
    store = MappingStore()
    store.load_activities([ActivityRow("A1000", "1", "1.1", "Concrete")])
    store.load_boq([BOQRow("BOQ|2", "BOQ", 2, "1", "", "", "Concrete", 1000, "BOQ-ONE")])
    store.selected_activity_ids = {"A1000"}
    store.selected_boq_ids = {"BOQ|2"}
    store.map_selected(75)
    source = WorkingTreeScheduleSource(
        _full_rebuild_fixture(tmp_path / "seed.xlsx"),
        list(store.working_tree_nodes()),
        {"A1000": 1000},
    )
    created = tmp_path / "created.xlsx"
    WorkbookGenerationService().generate(source, created, amounts={"A1000": 1000})
    wb = load_workbook(created)
    try:
        # Use the real persistent input adapter, then make the requirement explicit.
        PaymentInputWorkbook().embed(wb)
        payment = wb["Payment Input"]
        for row in range(8, payment.max_row + 1):
            if payment.cell(row, 1).value == "ACT":
                for col in range(5, payment.max_column + 1):
                    payment.cell(row, col).value = None
                if populated_payment:
                    payment.cell(row, 5, 0.5)
        assert len(wb["main"]._charts) == len(wb["main_monthly"]._charts) == 1
        wb.save(created)
    finally:
        wb.close()
    return created, store


@pytest.mark.parametrize("populated_payment", [False, True])
def test_mapping_final_output_restores_overlay_dependencies(tmp_path: Path, populated_payment: bool):
    """Assert the delivered file, including the real trailing Payment save when used."""
    source, store = _created_mapping_input(tmp_path, populated_payment)
    output = tmp_path / "mapped.xlsx"
    events = []
    result = WorkbookExportService().export(
        source, output, store,
        progress_callback=lambda step, message, complete: events.append((step, complete)),
    )
    assert result.validation.allocated_amount == 750
    assert result.validation.remaining_amount == 250
    assert result.mapping_rows_written == 1
    assert (("payment", True) in events) is populated_payment

    wb = load_workbook(output, data_only=False)
    try:
        if populated_payment:
            assert wb["Payment"]._images
            assert not wb["Payment"]._charts
        data = wb["Dashboard_Data"]
        assert len(wb["Dashboard"]._charts) == 1
        assert len(wb["main"]._charts) == 1
        assert len(wb["main_monthly"]._charts) == 1

        # Display-only X columns remain visible but never become reporting points.
        weekly = [datetime(2026, 3, 6), datetime(2026, 3, 13)]
        monthly = [datetime(2026, 3, 13)]
        assert [data.cell(r, 10).value for r in (2, 3, 4)] == weekly + [None]
        assert [data.cell(r, 11).value for r in (2, 3)] == monthly + [None]
        for name, date_col, value_cols, dates, list_col in (
            ("main", "T", ("U", "V", "W"), weekly, "J"),
            ("main_monthly", "X", ("Y", "Z", "AA"), monthly, "K"),
        ):
            ws = wb[name]
            chart = ws._charts[0]
            last = len(dates) + 2  # existing chart-only leading zero point
            assert len(chart.series) == 3
            assert [data[f"{date_col}{r}"].value for r in range(3, last + 1)] == dates
            assert data[f"{date_col}2"].value < dates[0]
            assert data[f"{date_col}{last + 1}"].value is None
            assert data[f"{value_cols[0]}2"].value == 0
            assert data[f"{value_cols[1]}2"].value == 0
            for series, col in zip(chart.series, value_cols):
                assert series.cat.numRef.f == f"'Dashboard_Data'!${date_col}$2:${date_col}${last}"
                assert series.val.numRef.f == f"'Dashboard_Data'!${col}$2:${col}${last}"
                for ref in (series.cat.numRef.f, series.val.numRef.f):
                    sheet, (c1, r1, c2, r2) = range_to_tuple(ref)
                    assert all(wb[sheet].cell(r, c).value is not None
                               for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
            assert chart.series[2].tx.v == "Cutoff"
            assert chart.series[2].errBars.errBarType == "minus"
            assert chart.display_blanks == "gap"
            reporting_cols = [c.column for c in ws[3]
                              if isinstance(c.value, str) and c.value.startswith(("W", "M"))]
            assert any(c.value == "X" for c in ws[3])
            assert chart.anchor.editAs == "twoCell"
            assert chart.anchor._from.col == min(reporting_cols) - 1
            assert chart.anchor.to.col == max(reporting_cols)
            assert chart.anchor._from.row == 4
            cutoff_row = next(r for r in range(1, ws.max_row + 1)
                              if ws.cell(r, 12).value == "Cutoff Date")
            assert chart.anchor.to.row == cutoff_row
            control = ws.cell(cutoff_row, 13)
            assert control.value == dates[-1]
            assert not control.protection.locked
            validations = [dv for dv in ws.data_validations.dataValidation
                           if control.coordinate in dv.sqref]
            assert len(validations) == 1
            assert validations[0].formula1 == f'=INDIRECT("Dashboard_Data!${list_col}$2:${list_col}${len(dates) + 1}")'
            cutoff_name = "PS_WEEKLY_OVERLAY_CUTOFF" if name == "main" else "PS_MONTHLY_OVERLAY_CUTOFF"
            assert wb.defined_names[cutoff_name].attr_text == f"'{name}'!$M${cutoff_row}"

        # Series helpers resolve to the existing reporting/visible-Actual sources.
        for target, source_col in (("U", "B"), ("V", "P"), ("W", "R"),
                                   ("Y", "E"), ("Z", "Q"), ("AA", "S")):
            assert data[f"{source_col}2"].value is not None
            assert f"{source_col}2" in data[f"{target}3"].value
        assert "PS_WEEKLY_OVERLAY_CUTOFF" in data["P2"].value
        assert "PS_MONTHLY_OVERLAY_CUTOFF" in data["Q2"].value

        # Reconciliation and mapped business values survive the final renderer/save.
        main = wb["main"]
        headers = {str(c.value).lower(): c.column for c in main[4]}
        activity_row = next(r for r in range(5, main.max_row + 1)
                            if main.cell(r, headers["activity id"]).value == "A1000"
                            and main.cell(r, headers["p/a"]).value == "P")
        assert main.cell(activity_row, headers["amount"]).value == 750
        assert wb["Mapping Summary"]["B14"].value == 750
        assert wb["Mapping Summary"]["B15"].value == 250
        assert wb["BOQ Activity Mapping"].max_row == 2
        assert wb.calculation.calcMode == "manual"
        assert wb.calculation.fullCalcOnLoad is False
        assert wb.calculation.forceFullCalc is False
        assert wb.calculation.calcOnSave is True
    finally:
        wb.close()
