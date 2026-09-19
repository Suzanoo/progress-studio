"""Exercise real workspace commands/services with headless Tk variables.

Only widget drawing, file dialogs and OS launch are substituted. Workbook reads,
worker threads, queued completion and finance publication use production code.
Desktop visual/Excel acceptance remains a separate Product Owner gate.
"""
from datetime import date
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
import tkinter as tk
from openpyxl import Workbook, load_workbook

from progress_studio.presentation.gui import finance
from progress_studio.presentation.gui.app import ProgressStudioDesktopApp
from progress_studio.infrastructure.excel.finance_input_workbook import read_finance_inputs
from progress_studio.services.financial_forecast_deriver import FinancialForecastDeriver
from tests.unit.finance.test_finance_workbook import sample
from tests.fixtures.finance_workbook_probe import package_issues


class Control:
    def __init__(self):
        self.options = {}

    def configure(self, **options):
        self.options.update(options)


@pytest.fixture
def workspace(monkeypatch):
    interpreter = tk.Tcl()
    variable = tk.StringVar
    monkeypatch.setattr(finance.tk, "StringVar", lambda **kw: variable(master=interpreter, **kw))
    monkeypatch.setattr(finance.ttk.Frame, "__init__", lambda *a, **kw: None)

    def build(frame):
        for name in ("path_entry", "browse_button", "check_button", "capacity_entry", "generate_button", "open_button"):
            setattr(frame, name, Control())
        frame.initial_entries = [Control() for _ in range(3)]
        frame._update_controls()

    monkeypatch.setattr(finance.FinanceFrame, "_build_ui", build)
    frame = finance.FinanceFrame(None)
    frame.callbacks = []
    monkeypatch.setattr(frame, "after", lambda delay, callback: frame.callbacks.append(callback))
    frame.warning = Mock()
    frame.error = Mock()
    monkeypatch.setattr(finance.messagebox, "showwarning", frame.warning)
    monkeypatch.setattr(finance.messagebox, "showerror", frame.error)
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: "")
    monkeypatch.setattr(finance.filedialog, "askopenfilename", lambda **kw: "")
    yield frame
    if frame._worker:
        frame._worker.join(timeout=30)
        assert not frame._worker.is_alive()


def finish(frame):
    frame._worker.join(timeout=30)
    assert not frame._worker.is_alive()
    callback = frame.callbacks.pop(0)
    callback()
    assert not frame.busy


def check(frame, path):
    frame.workbook_var.set(str(path))
    frame._check()
    finish(frame)


@pytest.fixture
def source(tmp_path):
    path = tmp_path / "progress.xlsx"
    w = Workbook()
    w.active.title = "main"
    w.active["A1"] = "Existing project"
    w.save(path)
    w.close()
    return path


def prepare_creation(frame, source):
    check(frame, source)
    frame.opening_var.set("2026-01-01")
    frame.actuals_var.set("2026-01-31")
    frame.currency_var.set("THB")
    frame.capacity_var.set("4")


@pytest.mark.smoke
def test_create_through_workspace_keeps_source_and_opens_new_result(workspace, source, tmp_path, monkeypatch):
    original = source.read_bytes()
    prepare_creation(workspace, source)
    assert workspace.generate_button.options["text"] == "Create Finance Workbook"
    assert workspace.initial_entries[0].options["state"] == "normal"
    output = tmp_path / "finance.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(output))
    workspace._generate()
    assert workspace.busy
    assert workspace.browse_button.options["state"] == "disabled"
    workspace._generate()  # duplicate activation cannot start another worker
    assert len(workspace.callbacks) == 1
    finish(workspace)
    workspace.error.assert_not_called()
    assert source.read_bytes() == original
    assert workspace.workbook_var.get() == str(output)
    assert workspace._analysis.existing_finance
    assert workspace.initial_entries[0].options["state"] == "disabled"
    assert workspace.open_button.options["state"] == "normal"
    assert not package_issues(output)
    w = load_workbook(output)
    inputs = read_finance_inputs(w)
    assert inputs.opening_balance == 0
    assert not inputs.actuals_complete and not inputs.receivables_complete and not inputs.cash_out_complete
    assert w["Finance Input"]["B8"].value == 30
    assert w["Finance Input"]["B9"].value == .05
    assert w["Cash Flow"].sheet_state == "visible"
    assert w["Finance_Data"].sheet_state == "hidden"
    assert not w["Finance Input"]["B6"].protection.locked
    w.close()
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.read("xl/worksheets/sheet1.xml") == after.read("xl/worksheets/sheet1.xml")
    opener = Mock()
    monkeypatch.setattr(finance.os, "name", "posix")
    monkeypatch.setattr(finance.sys, "platform", "darwin")
    monkeypatch.setattr(finance.os, "spawnlp", opener, raising=False)
    workspace._open_result()
    opener.assert_called_once_with(finance.os.P_NOWAIT, "open", "open", str(output))


def test_refresh_preserves_saved_inputs_zero_override_notes_and_expands(workspace, tmp_path, monkeypatch):
    source = tmp_path / "edited.xlsx"
    w = sample()
    s = w["Finance Input"]
    s["U28"] = "cert:net"
    s["V28"] = 0
    s["Y28"] = "Cancel remaining"
    s["Z28"] = "Keep this note"
    expected = read_finance_inputs(w)
    w.save(source)
    w.close()
    original = source.read_bytes()
    check(workspace, source)
    assert workspace.generate_button.options["text"] == "Refresh Finance Workbook"
    assert workspace.opening_var.get() == "2026-01-01"
    workspace.opening_var.set("invalid disabled setting")
    workspace.currency_var.set("IGNORED")
    workspace.capacity_var.set("6")
    output = tmp_path / "refreshed.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(output))
    workspace._generate()
    finish(workspace)
    workspace.error.assert_not_called()
    assert source.read_bytes() == original
    assert workspace._analysis.capacity_rows == 6
    w = load_workbook(output)
    assert read_finance_inputs(w) == expected
    assert w["Finance Input"]["Z28"].value == "Keep this note"
    assert w["Finance Input"]["V28"].value == 0
    assert not w["Finance Input"]["V30"].protection.locked
    result = FinancialForecastDeriver().derive(read_finance_inputs(w))
    assert result.peak_funding_requirement == 50
    assert result.minimum_balance_date == date(2026, 2, 5)
    assert result.remaining_receivable == 125
    w.close()
    assert not package_issues(output)


@pytest.mark.parametrize("field,value", [
    ("opening_var", ""), ("opening_var", "2026-02-30"),
    ("actuals_var", "2025-12-31"), ("currency_var", " "),
    ("capacity_var", "0"), ("capacity_var", "2001"), ("capacity_var", "1.5"),
])
def test_invalid_creation_fields_do_not_offer_save(workspace, source, monkeypatch, field, value):
    prepare_creation(workspace, source)
    getattr(workspace, field).set(value)
    save = Mock()
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", save)
    workspace._generate()
    save.assert_not_called()
    workspace.warning.assert_called_once()
    assert not workspace.busy


@pytest.mark.parametrize("choice", ["cancel", "source", "existing", "xlsm"])
def test_save_cancel_or_unsafe_destination_never_writes(workspace, source, tmp_path, monkeypatch, choice):
    prepare_creation(workspace, source)
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"do not overwrite")
    paths = {"cancel": "", "source": str(source), "existing": str(existing), "xlsm": str(tmp_path / "bad.xlsm")}
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: paths[choice])
    generate = Mock()
    monkeypatch.setattr(workspace.service, "generate", generate)
    workspace._generate()
    generate.assert_not_called()
    assert existing.read_bytes() == b"do not overwrite"
    assert not workspace.busy


def test_typed_path_change_invalidates_previous_check(workspace, source, tmp_path):
    check(workspace, source)
    workspace.workbook_var.set(str(tmp_path / "different.xlsx"))
    assert workspace._analysis is None
    assert workspace.generate_button.options["state"] == "disabled"
    assert workspace.open_button.options["state"] == "disabled"
    workspace._generate()
    workspace.warning.assert_called_once()


def test_browse_and_cancel(workspace, source, monkeypatch):
    workspace._browse()
    assert workspace._worker is None
    monkeypatch.setattr(finance.filedialog, "askopenfilename", lambda **kw: str(source))
    workspace._browse()
    finish(workspace)
    assert workspace._analysis.workbook == source


def test_saved_file_is_revalidated_after_check_and_error_delivered_on_ui_thread(workspace, source, tmp_path, monkeypatch):
    prepare_creation(workspace, source)
    w = load_workbook(source)
    w.create_sheet("Cash Flow")["A1"] = "User-owned sheet"
    w.save(source)
    w.close()
    output = tmp_path / "failure.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(output))
    workspace._generate()
    workspace._worker.join(timeout=30)
    workspace.error.assert_not_called()  # worker must not display dialogs
    finish(workspace)
    workspace.error.assert_called_once()
    assert "collision" in str(workspace.error.call_args)
    assert not output.exists()
    assert not list(tmp_path.glob(".finance.*"))
    assert workspace._analysis is None
    assert workspace.check_button.options["state"] == "normal"


def test_generation_failure_restores_controls(workspace, source, tmp_path, monkeypatch):
    prepare_creation(workspace, source)
    output = tmp_path / "locked.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(output))
    monkeypatch.setattr(workspace.service, "generate", Mock(side_effect=PermissionError("Workbook is locked")))
    workspace._generate()
    finish(workspace)
    assert "Workbook is locked" in str(workspace.error.call_args)
    assert workspace.open_button.options["state"] == "disabled"
    assert workspace.check_button.options["state"] == "normal"
    assert not output.exists()


def test_missing_source_and_missing_result_are_visible_errors(workspace, tmp_path):
    check(workspace, tmp_path / "missing.xlsx")
    workspace.error.assert_called_once()
    workspace.error.reset_mock()
    workspace._output_path = tmp_path / "moved.xlsx"
    workspace._open_result()
    assert "Could not open workbook" in str(workspace.error.call_args)


def test_shell_routes_finance_and_prevents_close_during_work(monkeypatch):
    assert ("finance", "¤", "Finance") in ProgressStudioDesktopApp.WORKSPACES
    frame = Mock()
    shell = SimpleNamespace(
        workspace_frames={"finance": frame}, current_workspace="home",
        WORKSPACES=ProgressStudioDesktopApp.WORKSPACES,
        workspace_title_var=Mock(), sidebar_buttons={"finance": Mock()}, command_bar=Mock(),
        finance_workspace=SimpleNamespace(busy=True), _save_layout_preferences=Mock(), destroy=Mock(),
    )
    ProgressStudioDesktopApp._show_workspace(shell, "finance")
    frame.tkraise.assert_called_once()
    assert shell.current_workspace == "finance"
    shell.command_bar.grid_remove.assert_called_once()
    warning = Mock()
    monkeypatch.setattr(finance.messagebox, "showwarning", warning)
    ProgressStudioDesktopApp._close_application(shell)
    warning.assert_called_once()
    shell.destroy.assert_not_called()
    shell.finance_workspace.busy = False
    ProgressStudioDesktopApp._close_application(shell)
    shell.destroy.assert_called_once()


def test_existing_capacity_cannot_shrink(workspace, tmp_path, monkeypatch):
    source = tmp_path / "existing.xlsx"
    w = sample()
    w.save(source)
    w.close()
    check(workspace, source)
    workspace.capacity_var.set("3")
    save = Mock()
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", save)
    workspace._generate()
    save.assert_not_called()
    assert "cannot be reduced" in str(workspace.warning.call_args)


def test_finance_added_after_check_requires_new_check(workspace, source, tmp_path, monkeypatch):
    prepare_creation(workspace, source)
    w = sample()
    w.save(source)
    w.close()
    output = tmp_path / "output.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(output))
    workspace._generate()
    finish(workspace)
    assert "changed since checking" in str(workspace.error.call_args)
    assert not output.exists()


def test_workspace_finance_round_trip_through_real_mapping_and_rebuild(workspace, tmp_path, monkeypatch):
    from tests.fixtures.finance_workbook_probe import create_seed, _sheet_parts
    from progress_studio.services.rebuild_service import WorkbookRebuildEngine
    from progress_studio.services.workbook_export_service import WorkbookExportService
    from progress_studio.services.mapping_store import MappingStore
    from progress_studio.domain.mapping_models import ActivityRow, BOQRow

    source = create_seed(tmp_path / "project")
    prepare_creation(workspace, source)
    finance_path = tmp_path / "project_finance.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(finance_path))
    workspace._generate()
    finish(workspace)
    workspace.error.assert_not_called()
    with ZipFile(source) as before, ZipFile(finance_path) as after:
        original = {name: before.read(name) for name in before.namelist()}
        for part in _sheet_parts(original).values():
            assert after.read(part) == original[part]
        for part in original:
            if part.startswith(("xl/charts/", "xl/drawings/")):
                assert after.read(part) == original[part]
    edited = load_workbook(finance_path)
    s = edited["Finance Input"]
    for col, value in enumerate(["supplier", 100, "out", date(2026, 2, 5), "operating", None, False, "manual"], 10):
        s.cell(25, col, value)
    s["S25"] = "Saved in Excel"
    s["U28"] = "supplier"
    s["V28"] = 0
    s["Y28"] = "Cancelled remaining"
    expected = read_finance_inputs(edited)
    edited.save(finance_path)
    edited.close()

    store = MappingStore()
    store.load_activities([ActivityRow("A1000", "1", "1.1", "Concrete")])
    store.load_boq([BOQRow("BOQ|2", "BOQ", 2, "1", "", "", "Concrete", 1000, "BOQ-ONE")])
    store.selected_activity_ids = {"A1000"}
    store.selected_boq_ids = {"BOQ|2"}
    store.map_selected(100)
    mapped = tmp_path / "mapped.xlsx"
    WorkbookExportService().export(finance_path, mapped, store)
    rebuilt = tmp_path / "rebuilt.xlsx"
    WorkbookRebuildEngine().rebuild_live_progress(mapped, rebuilt)
    check(workspace, rebuilt)
    final = tmp_path / "final.xlsx"
    monkeypatch.setattr(finance.filedialog, "asksaveasfilename", lambda **kw: str(final))
    workspace._generate()
    finish(workspace)
    workspace.error.assert_not_called()
    result = load_workbook(final)
    assert read_finance_inputs(result) == expected
    assert result["Finance Input"]["S25"].value == "Saved in Excel"
    assert not result["Finance Input"]["V28"].protection.locked
    assert result["Finance_Data"].sheet_state == "hidden"
    result.close()
    assert not package_issues(final)
