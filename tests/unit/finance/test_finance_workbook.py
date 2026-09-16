from copy import copy
from datetime import date
from decimal import Decimal
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from progress_studio.infrastructure.excel.finance_input_workbook import (
    create_inputs, read_finance_inputs, preserve_finance_inputs, check_schema, INPUT,
)
from progress_studio.infrastructure.excel.financial_forecast_workbook import build_finance_view
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.services.financial_forecast_workbook_service import FinancialForecastWorkbookService
from progress_studio.services.financial_forecast_deriver import FinancialForecastDeriver, FinanceValidationError


def sample():
    w=Workbook();w.active.title='main'
    ws=create_inputs(w,opening_date=date(2026,1,1),actuals_through=date(2026,1,31),currency='THB',capacity_rows=4)
    for r in range(10,14):ws.cell(r,2,True)
    ws['B6']=10
    for r,values in enumerate([
        ['pay',100,'out',date(2026,2,5),'operating',None,False,'manual'],
        ['invoice',120,'in',date(2026,2,25),'operating',None,False,'manual'],
        ['cert',100,'in',date(2026,1,31),'operating',None,False,'certificate',date(2026,6,1)],
    ],25):
        for c,value in enumerate(values,10):ws.cell(r,c,value)
    for c,value in enumerate(['receipt',date(2026,1,20),40,'in','operating','cert:net',None,'Paid cash'],1):ws.cell(25,c,value)
    return w


def test_reader_preserves_exact_dates_and_ff1_meaning():
    w=sample();result=FinancialForecastDeriver().derive(read_finance_inputs(w))
    assert result.peak_funding_requirement==50
    assert result.minimum_balance_date==date(2026,2,5)
    assert result.remaining_receivable==180
    assert result.monthly_points[-1].closing_balance==130
    assert next(i for i in result.items if i.item.item_id=='cert:net').effective_date==date(2026,3,2)


def test_snapshot_override_zero_and_clear_date_are_distinct():
    w=sample();s=w[INPUT];s['U25']='cert:net';s['V25']=0;s['Y25']='Cancel remaining'
    result=FinancialForecastDeriver().derive(read_finance_inputs(w));assert result.remaining_receivable==125
    s['V25']=None;s['X25']=True
    result=FinancialForecastDeriver().derive(read_finance_inputs(w))
    assert result.remaining_receivable==180
    assert 'undated:cert:net' in result.issues
    s['X25']=False;s['W25']=date(2026,4,4)
    result=FinancialForecastDeriver().derive(read_finance_inputs(w))
    assert next(i for i in result.items if i.item.item_id=='cert:net').effective_date==date(2026,4,4)


@pytest.mark.parametrize('cell,value',[
 ('A26','receipt'),('B25',date(2026,1,20)),('C25',-1),('F25','unknown'),
 ('J28','CERT:NET'),('K25','text'),('M25','2026-02-05'),('O25',date(2027,1,1)),
 ('U25','missing'),('B8',1.5),('B9',1.1),('B10','yes'),('B6',1e14),('A25','wild*card'),
])
def test_invalid_inputs_are_rejected(cell,value):
    w=sample();s=w[INPUT];s[cell]=value
    if cell=='A26':
        for c in range(2,8):s.cell(26,c,s.cell(25,c).value)
    elif cell=='B25':s['B25']='not a date'
    elif cell=='J28':
        for c in range(11,18):s.cell(28,c,s.cell(26,c).value)
    with pytest.raises((FinanceValidationError,ValueError)):read_finance_inputs(w)


def test_notes_last_row_and_inputs_survive_view_refresh_and_finalization():
    w=sample();s=w[INPUT];s['H25']='Keep note';s['U28']='invoice';s['V28']=0;s['Y28']='Cancelled'
    before=[tuple(c.value for c in row) for row in s]
    build_finance_view(w);finalize_workbook(w);build_finance_view(w)
    assert before==[tuple(c.value for c in row) for row in s]
    assert s.sheet_state=='visible' and w['Cash Flow'].sheet_state=='visible'
    assert w['Finance_Data'].sheet_state=='hidden'
    assert not s['V28'].protection.locked and s['U24'].protection.locked
    assert w['Cash Flow']['B7'].protection.locked
    assert w.calculation.calcMode=='manual'
    assert w['Cash Flow']['B7'].data_type=='f'


def test_mapping_transfer_keeps_values_notes_overrides_and_zero():
    source=sample();source[INPUT]['U28']='invoice';source[INPUT]['V28']=0;source[INPUT]['Y28']='Keep this';source[INPUT]['Z28']='Note'
    target=Workbook();target.active.title='main'
    preserve_finance_inputs(source,target)
    assert read_finance_inputs(source)==read_finance_inputs(target)
    assert target[INPUT]['Z28'].value=='Note'
    assert target['Cash Flow']['B7'].data_type=='f'


def test_collision_and_overflow_fail_without_guessing():
    w=sample();w[INPUT]['A29']='outside'
    with pytest.raises(FinanceValidationError,match='capacity'):check_schema(w)
    w[INPUT]['A29']=None;w[INPUT]['A1']='FF-0 DISPOSABLE PROBE v1'
    with pytest.raises(FinanceValidationError,match='schema'):check_schema(w)


def test_extension_publisher_preserves_unowned_parts_exactly(tmp_path):
    w=Workbook();s=w.active;s.title='main';s.append([1,2]);s.append([2,3])
    chart=LineChart();chart.add_data(Reference(s,min_col=2,min_row=1,max_row=2));s.add_chart(chart,'D2')
    source=tmp_path/'source.xlsx';w.save(source)
    out=tmp_path/'finance.xlsx';again=tmp_path/'again.xlsx'
    service=FinancialForecastWorkbookService()
    service.generate(source,out,opening_date=date(2026,1,1),actuals_through=date(2026,1,31),currency='THB',capacity_rows=4)
    service.generate(out,again)
    with ZipFile(source) as a,ZipFile(again) as b:
        allowed={'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','xl/styles.xml'}
        for name in a.namelist():
            if name not in allowed:assert a.read(name)==b.read(name),name
    check=load_workbook(again);assert len(check['main']._charts)==1;read_finance_inputs(check);check.close()
    assert source.read_bytes()!=again.read_bytes()
    with pytest.raises(ValueError):service.generate(source,source)
    with pytest.raises(ValueError):service.generate(source,out)


def test_refresh_after_excel_style_table_renumbering(tmp_path):
    w=sample();build_finance_view(w);source=tmp_path/'edited.xlsx';w.save(source)
    output=tmp_path/'out.xlsx';FinancialForecastWorkbookService().generate(source,output)
    with ZipFile(output) as z:
        tables=[ET.fromstring(z.read(n)).get('name') for n in z.namelist() if n.startswith('xl/tables/') and n.endswith('.xml')]
        assert len(tables)==len(set(tables))==3
    result=load_workbook(output);assert read_finance_inputs(result)==read_finance_inputs(w)
