import math
import pytest
from openpyxl import Workbook, load_workbook
from progress_studio.app.desktop import DesktopRunOptions, DesktopRunner
from progress_studio.app.pipeline import Pipeline
from progress_studio.config import SETTINGS
from progress_studio.presentation.cli import CommandLineInterface
from progress_studio.domain import Activity
from progress_studio.services.weighting import assign_dummy_weights, validate_weight_basis
from progress_studio.infrastructure.schedule_xml.xml_reader import ScheduleXmlReader
from progress_studio.infrastructure.excel.weight_basis import set_creation_basis, creation_basis, apply_weight_labels, uses_dummy_weights


def activity(aid='A1', hours=8, milestone=False, summary=False):
    return Activity(0,None,None,aid,aid,'1',1,summary,None,None,None,None,None,None,None,999,hours,milestone)


def test_equal_ignores_source_amount_duration_and_summary():
    rows=[activity(),activity('A2',40),activity('M',0,True),activity('W',None,summary=True)]
    assign_dummy_weights(rows,'equal')
    assert [r.amount for r in rows[:3]]==[1,1,0]
    assert rows[3].amount==999


def test_duration_uses_source_hours_and_preserves_milestone():
    rows=[activity(),activity('A2',40),activity('M',0,True)]
    assign_dummy_weights(rows,'duration')
    assert [r.amount for r in rows]==[8,40,0]
    assert len(rows)==3


@pytest.mark.parametrize('hours',[None,0,-1,math.inf,math.nan])
def test_invalid_duration_is_actionable_and_atomic(hours):
    rows=[activity(),activity('BAD',hours)]
    with pytest.raises(ValueError,match='BAD'):
        assign_dummy_weights(rows,'duration')
    assert rows[0].amount==999


@pytest.mark.parametrize('basis',['equal','duration'])
def test_all_milestone_total_rejected(basis):
    with pytest.raises(ValueError,match='total'):
        assign_dummy_weights([activity('M',0,True)],basis)


@pytest.mark.parametrize('value,expected',[('PT8H30M0S',8.5),('PT90M',1.5),('PT30S',1/120),('PT0H0M0S',0),('P1D',None),('PT',None),('-PT2H',None),('bad',None),('',None)])
def test_msp_duration_parser(value,expected):
    assert ScheduleXmlReader._duration_hours(value)==expected


def test_desktop_default_and_amount_unavailable(tmp_path):
    xml=tmp_path/'a.xml';xml.write_text('<Project/>')
    runner=DesktopRunner(lambda _:Pipeline([]))
    assert runner.run(DesktopRunOptions(xml,'5')).weight_basis=='equal'
    assert runner.run(DesktopRunOptions(xml,'5',weight_basis='duration')).weight_basis=='duration'
    with pytest.raises(ValueError,match='unavailable'):
        runner.run(DesktopRunOptions(xml,'5',weight_basis='amount'))
    with pytest.raises(ValueError,match='basis'):
        validate_weight_basis('bogus')
    cli=CommandLineInterface(SETTINGS)
    assert cli.parse([]).weight_basis=='equal'
    assert cli.parse(['--weight-basis','duration']).weight_basis=='duration'


def test_basis_save_reopen_and_mapping_label(tmp_path):
    wb=Workbook();wb.active.title='main';wb['main']['I4']='Amount'
    wb.create_sheet('README');wb.create_sheet('Dashboard')['J6']='Project Value'
    wb.create_sheet('Payment-Breakdown')['A1']='Payment Breakdown'
    wb['Payment-Breakdown']['D4']='Amount'
    wb['main']['Q4']='XML Amount'
    set_creation_basis(wb,'duration');apply_weight_labels(wb)
    assert wb['Dashboard']['J6'].value=='Total weight'
    assert 'dummy-weighted' in wb['Payment-Breakdown']['A1'].value
    assert 'not Contract Value' in wb['Payment-Breakdown']['D4'].comment.text
    assert 'not Contract Value' in wb['main']['Q4'].comment.text
    assert 'not Contract Value' in wb['README']['B4'].value
    path=tmp_path/'basis.xlsx';wb.save(path);wb.close()
    wb=load_workbook(path)
    assert creation_basis(wb)=='duration'
    wb.create_sheet('BOQ Activity Mapping');apply_weight_labels(wb)
    assert not uses_dummy_weights(wb)
    assert 'Current amounts: BOQ Mapping' in wb['README']['B4'].value
    assert wb['Dashboard']['J6'].value=='Project Value'
    assert wb['Payment-Breakdown']['A1'].value=='Payment Breakdown'
    assert 'BOQ Mapping' in wb['main']['I4'].comment.text
    wb.close()
