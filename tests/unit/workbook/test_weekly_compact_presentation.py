from copy import deepcopy
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.chart.axis import TextAxis

from scripts.generate_weekly_compact_acceptance import build_source
from progress_studio.infrastructure.excel.main_dataset_workbook_adapter import main_dataset_from_workbook
from progress_studio.infrastructure.excel.live_dashboard_workbook import build_live_dashboard
from progress_studio.infrastructure.excel.live_monthly_workbook import build_live_monthly_view
from progress_studio.services.monthly_cache_deriver import MonthlyCacheDeriver
from progress_studio.infrastructure.excel.traditional_overlay_workbook import build_traditional_overlays
from progress_studio.infrastructure.excel.weekly_compact_workbook import build_weekly_compact
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook


def fingerprint(ws):
    return (
        [(c.coordinate, c.value, tuple(c._style or [0]*9), str(c.protection))
         for row in ws for c in row],
        sorted(str(r) for r in ws.merged_cells.ranges),
        [(k, str(v)) for k,v in ws.column_dimensions.items()],
        [(k, str(v)) for k,v in ws.row_dimensions.items()],
        [(s.val.numRef.f, str(s.cat), str(s.dLbls), str(s.marker))
         for ch in ws._charts for s in ch.series],
        str(ws.data_validations), str(ws.protection), ws.freeze_panes,
        [(str(r), [str(v) for v in rules]) for r,rules in ws.conditional_formatting._cf_rules.items()],
    )


@pytest.fixture(scope="module")
def finalized(tmp_path_factory):
    folder = tmp_path_factory.mktemp("compact")
    wb = build_source(folder)
    ds = main_dataset_from_workbook(wb)
    build_live_monthly_view(wb, ds, MonthlyCacheDeriver().derive(ds))
    build_live_dashboard(wb, ds)
    build_traditional_overlays(wb, ds)
    return wb, ds, folder


def test_projection_preserves_both_existing_sheets_and_canonical_helpers(finalized):
    wb, ds, _ = finalized
    before = {n:fingerprint(wb[n]) for n in ("main", "main_monthly", "Dashboard_Data")}
    assert build_weekly_compact(wb, ds)
    assert before == {n:fingerprint(wb[n]) for n in before}
    source, compact = wb["main"], wb["Weekly Compact"]
    dates = [c.column for c in source[4] if isinstance(c.value, datetime)]
    assert len(ds.periods) == 310
    assert len(dates) == 316  # original three-week margins each side
    assert [compact.cell(4,c).value for c in dates] == [source.cell(4,c).value for c in dates]
    assert [compact.cell(3,c).value for c in dates] == [source.cell(3,c).value for c in dates]
    for row in ds.activities:
        for r in (row.row_number,row.row_number+1):
            for col in dates:
                coordinate = source.cell(r,col).coordinate
                assert compact.cell(r,col).value == f'=IF(\'main\'!{coordinate}="","",\'main\'!{coordinate})'
    assert compact.freeze_panes == source.freeze_panes
    assert [d.outlineLevel for d in compact.row_dimensions.values()] == [d.outlineLevel for d in source.row_dimensions.values()]


@pytest.mark.parametrize("year,month,count", [(2026,10,5),(2026,11,4),(2024,2,4),(2023,2,4),(2025,12,4),(2026,1,5)])
def test_month_groups_keep_weekly_geometry(finalized, year, month, count):
    wb, ds, _ = finalized
    build_weekly_compact(wb,ds)
    s, c = wb["main"], wb["Weekly Compact"]
    cols = [cell.column for cell in s[4] if isinstance(cell.value,datetime) and (cell.value.year,cell.value.month)==(year,month)]
    assert len(cols) == count
    assert any(r.min_row==2 and r.min_col==cols[0] and r.max_col==cols[-1] for r in c.merged_cells.ranges)
    assert c.cell(5,cols[0]).border.left.style in ("thin","medium")
    for col in cols[1:]:
        assert c.cell(5,col).border.left.style is None
        assert c.cell(5,col-1).border.right.style is None


def test_hidden_labels_uniform_width_partial_months_and_year_boundary(finalized):
    wb, ds, _ = finalized
    build_weekly_compact(wb,ds)
    s,c=wb["main"],wb["Weekly Compact"]
    dates=[x for x in s[4] if isinstance(x.value,datetime)]
    for cell in dates:
        assert c.column_dimensions[cell.column_letter].width == 2.5
        assert not c.column_dimensions[cell.column_letter].hidden
        assert c.cell(3,cell.column).number_format == ";;;"
        assert c.cell(4,cell.column).number_format == ";;;"
        assert c.cell(5,cell.column).number_format == ";;;"
    assert ds.periods[0].reporting_date == datetime(2022,1,14)
    assert ds.periods[-1].reporting_date == datetime(2027,12,17)
    jan=next(x.column for x in dates if x.value==datetime(2026,1,2))
    assert c.cell(5,jan).border.left.style == "medium"


def test_full_category_chart_same_values_cutoff_anchor_and_shared_state(finalized):
    wb,ds,_=finalized
    build_weekly_compact(wb,ds)
    original, compact=wb["main"]._charts[0],wb["Weekly Compact"]._charts[0]
    assert isinstance(compact.x_axis,TextAxis)
    assert compact.anchor == original.anchor
    for a,b in zip(original.series,compact.series):
        assert a.val.numRef.f == b.val.numRef.f
        assert a.cat == b.cat
        assert a.val.numRef.f.endswith("$312")  # 310 + leading point + header
    for s in compact.series[:2]:
        assert s.dLbls is None and s.marker.symbol is None and s.smooth is False
        assert b'val="none"' in ET.tostring(s.marker.to_tree())
    assert original.series[0].dLbls.showVal
    assert compact.series[2].errBars == original.series[2].errBars
    assert compact.series[2].dLbls == original.series[2].dLbls
    name=wb.defined_names["PS_WEEKLY_OVERLAY_CUTOFF"]
    assert next(name.destinations)[0] == "main"
    assert not wb["Weekly Compact"].data_validations.dataValidation
    assert "PS_WEEKLY_OVERLAY_CUTOFF" in wb["Dashboard_Data"]["P2"].value
    assert "PS_WEEKLY_OVERLAY_CUTOFF" in wb["Dashboard_Data"]["R2"].value


def test_policy_and_package_roundtrip(finalized):
    wb,ds,folder=finalized
    build_weekly_compact(wb,ds)
    finalize_workbook(wb,mode="live")
    ws=wb["Weekly Compact"]
    assert ws.sheet_state=="visible" and ws.protection.sheet
    assert all(c.protection.locked for row in ws for c in row if not isinstance(c,MergedCell))
    path=folder/"roundtrip.xlsx";wb.save(path)
    other=load_workbook(path)
    finalize_workbook(other,mode="live")
    other.save(folder/"resaved.xlsx")
    assert len(other["Weekly Compact"]._charts)==1
    chart=other["Weekly Compact"]._charts[0]
    assert chart.plot_area.graphicalProperties.noFill
    with ZipFile(folder/"resaved.xlsx") as z:
        charts=[ET.fromstring(z.read(n)) for n in z.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml")]
        ns={"c":"http://schemas.openxmlformats.org/drawingml/2006/chart"}
        assert any(r.find(".//c:catAx",ns) is not None and len(r.findall(".//c:ser",ns))==3 for r in charts)
    other.close()


def test_snapshot_freezes_grid_and_keeps_shared_overlay(finalized):
    wb,ds,_=finalized
    cached=deepcopy(wb["main"])
    # Stand-in cached result is distinct from its formula to expose accidental copy.
    col=ds.periods[0].column
    cached.cell(5,col,0.123)
    build_weekly_compact(wb,ds,snapshot=True,value_source=cached)
    assert wb["Weekly Compact"].cell(5,col).value==0.123
    for row in ds.activities:
        for p in ds.periods:
            assert wb["Weekly Compact"].cell(row.row_number,p.column).value==cached.cell(row.row_number,p.column).value
    assert wb["Weekly Compact"]._charts[0].series[1].val.numRef.f == wb["main"]._charts[0].series[1].val.numRef.f
