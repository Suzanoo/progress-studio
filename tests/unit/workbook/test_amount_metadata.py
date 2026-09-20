from openpyxl import Workbook,load_workbook
from progress_studio.domain.amount_field import AmountField
from progress_studio.infrastructure.excel.weight_basis import set_creation_basis,set_creation_field,creation_field,apply_weight_labels,uses_dummy_weights


def test_field_metadata_retains_literal_names_and_monetary_labels(tmp_path):
    w=Workbook();w.active.title='main';w['main']['A1']='Amount'
    w.create_sheet('README');w.create_sheet('Dashboard')['J6']='Project Value'
    field=AmountField('p6:udf:98','=A1+1','P6','Double')
    set_creation_basis(w,'amount');set_creation_field(w,field);apply_weight_labels(w)
    assert not uses_dummy_weights(w)
    assert 'monetary XML values' in w['main']['A1'].comment.text
    assert w['Dashboard']['J6'].value=='Project Value'
    path=tmp_path/'metadata.xlsx';w.save(path);w.close()
    w=load_workbook(path)
    assert creation_field(w)==field
    assert all(c.data_type!='f' for row in w['Info'] for c in row)
    w.close()
