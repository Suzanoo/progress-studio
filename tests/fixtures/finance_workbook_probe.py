"""FF-0 disposable integration probe. Not a production financial model.

Uses real Progress Studio owners; never patches their global policy or source.
Only the three probe sheets are authored here. Negative lifecycle findings are
reported as findings, not repaired or hidden behind passing assertions.
"""
from __future__ import annotations

from copy import copy
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import posixpath
import tempfile
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.packaging.relationship import Relationship
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.external_link.external import (
    ExternalLink, ExternalBook, ExternalSheetNames,
    ExternalSheetDataSet, ExternalSheetData, ExternalRow, ExternalCell,
)

from progress_studio.infrastructure.excel.dashboard_workbook import (
    _COLORS, _FONT, _date_axis_for_line_chart,
)
from progress_studio.infrastructure.excel.final_workbook_policy import finalize_workbook
from progress_studio.infrastructure.excel.progress_workbook import replace_defined_name
from progress_studio.infrastructure.excel.traditional_overlay_workbook import (
    reassert_traditional_overlay_transparency,
)
from progress_studio.infrastructure.excel.xlsx_package_preservation import restore_opaque_workbook_parts
from progress_studio.infrastructure.excel.xlsx_package_validator import validate_xlsx_tables

BASE_SHA = "c3c4d05a5679dbe710ede5b47b9c79681156aee2"
OWNED = ("Finance Input", "Cash Flow", "Finance_Data")
SCHEMA = "FF-0 DISPOSABLE PROBE v1"
FIRST_ROW, LAST_ROW = 20, 39
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/package/2006/relationships"
D = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _xml(node):
    return ET.tostring(node, encoding="unicode")


def _sheet_parts(parts):
    rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    targets = {r.get("Id"): posixpath.normpath(posixpath.join("xl", r.get("Target")))
               if not r.get("Target", "").startswith("/") else r.get("Target").lstrip("/")
               for r in rels}
    root = ET.fromstring(parts["xl/workbook.xml"])
    return {s.get("name"): targets[s.get(f"{{{D}}}id")]
            for s in root.find(f"{{{S}}}sheets")}


def package_issues(path):
    """Validate all internal relationship targets, IDs and part content types.

    Does not claim schema validation or that Excel will accept a package.
    """
    issues = []
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            issues.append("duplicate ZIP part")
        if archive.testzip():
            issues.append("ZIP CRC failure")
        parts = {n: archive.read(n) for n in names}
    ct = ET.fromstring(parts["[Content_Types].xml"])
    overrides = {n.get("PartName").lstrip("/") for n in ct if n.tag.endswith("Override")}
    defaults = {n.get("Extension") for n in ct if n.tag.endswith("Default")}
    for part in overrides:
        if part not in parts:
            issues.append(f"content type refers to missing part: {part}")
    for part in parts:
        if part != "[Content_Types].xml" and part not in overrides and part.rsplit(".", 1)[-1] not in defaults:
            issues.append(f"missing content type: {part}")
        if not part.endswith(".rels"):
            continue
        root = ET.fromstring(parts[part])
        ids = [r.get("Id") for r in root]
        if len(ids) != len(set(ids)):
            issues.append(f"duplicate relationship ID: {part}")
        owner = "" if part == "_rels/.rels" else part.replace("/_rels/", "/")[:-5]
        for rel in root:
            if rel.get("TargetMode") == "External":
                continue
            target = rel.get("Target", "")
            resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
            if resolved not in parts:
                issues.append(f"missing relationship target: {part} -> {resolved}")
        if owner and owner in parts and owner.endswith(".xml"):
            for element in ET.fromstring(parts[owner]).iter():
                for attr, value in element.attrib.items():
                    if attr.startswith("{" + D + "}") and value not in ids:
                        issues.append(f"unresolved relationship reference: {owner} {value}")
    try:
        validate_xlsx_tables(path)
    except ValueError as exc:
        issues.append(str(exc))
    return sorted(set(issues))


def snapshot(path):
    """Semantic sheet snapshots tolerate ZIP numbering, detect owner changes."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(path, keep_links=True)
    try:
        sheets = {}
        for ws in wb:
            styles = {}
            cells = {}
            for row in ws:
                for c in row:
                    if c.value is None:
                        continue
                    if c.style_id not in styles:
                        styles[c.style_id] = (c.number_format, c.protection.locked,
                                              _xml(c.font.to_tree()), _xml(c.fill.to_tree()))
                    cells[c.coordinate] = (c.value, *styles[c.style_id])
            sheets[ws.title] = {
                "cells": cells, "state": ws.sheet_state,
                "protection": _xml(ws.protection.to_tree()),
                "merges": sorted(str(r) for r in ws.merged_cells.ranges),
                "validations": _xml(ws.data_validations.to_tree()),
                "tables": sorted(_xml(t.to_tree()) for t in ws.tables.values()),
                "charts": [_xml(c.to_tree()) for c in ws._charts],
                "anchors": [_xml(c.anchor.to_tree()) for c in ws._charts],
                "images": [hashlib.sha256(i._data()).hexdigest() for i in ws._images],
            }
        return {"sheets": sheets, "names": {k: v.attr_text for k, v in wb.defined_names.items()},
                "calculation": _xml(wb.calculation.to_tree())}
    finally:
        wb.close()


def compare(before, after, names):
    issues = []
    for name in names:
        if name not in after["sheets"]:
            issues.append(f"{name}: sheet missing")
            continue
        for key, value in before["sheets"][name].items():
            if value != after["sheets"][name][key]:
                issues.append(f"{name}: {key} changed")
    return issues


def opaque_parts(path):
    with ZipFile(path) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()
                if n.startswith(("xl/drawings/", "xl/externalLinks/"))}


def create_seed(directory):
    """Reuse real Create/Mapping/Payment/EV fixture paths, then add sentinels."""
    from tests.integration.mapping.test_mapping_overlay_export import _created_mapping_input
    from progress_studio.services.workbook_export_service import WorkbookExportService
    from progress_studio.services.earned_value_rebuild_service import EarnedValueRebuildService
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    source, store = _created_mapping_input(directory, True)
    store.selected_activity_ids = {"A1000"}
    store.selected_boq_ids = {"BOQ|2"}
    store.map_selected(100)
    mapped = directory / "mapped.xlsx"
    WorkbookExportService().export(source, mapped, store)
    output = directory / "baseline.xlsx"
    EarnedValueRebuildService().generate(mapped, output)
    wb = load_workbook(output)
    try:
        notes = wb["User Notes"] if "User Notes" in wb.sheetnames else wb.create_sheet("User Notes")
        notes["A1"] = "FF-0 synthetic user note; preserve this content"
        notes["A2"] = "='[1]Probe Source'!A1"
        notes["B1"] = 11
        notes["B2"] = "=B1*2"
        link = ExternalLink(externalBook=ExternalBook(
            sheetNames=ExternalSheetNames(sheetName=["Probe Source"]),
            sheetDataSet=ExternalSheetDataSet(sheetData=[ExternalSheetData(
                sheetId=0, row=[ExternalRow(r=1, cell=[ExternalCell(r="A1", t="n", v="123")])])]), id="rId1"))
        link.file_link = Relationship(type="externalLinkPath", Target="file:///FF0-SYNTHETIC-NOT-AVAILABLE.xlsx", TargetMode="External")
        wb._external_links.append(link)
        replace_defined_name(wb, "FF0_User_Sentinel", "'User Notes'!$B$1")
        finalize_workbook(wb)
        # Expose the existing user note only in this synthetic fixture.
        notes.sheet_state = "visible"
        wb.save(output)
    finally:
        wb.close()
    return output


def _style(ws):
    ws.sheet_view.showGridLines = False
    for col in "ABCDEFGH":
        ws.column_dimensions[col].width = 19
    ws.column_dimensions["A"].width = 35
    ws.freeze_panes = "A20" if ws.title == "Finance Input" else "A6"
    for row in ws:
        for c in row:
            c.font = Font(name=_FONT, size=11, color=_COLORS["text"])
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor=_COLORS["navy"])
        c.font = Font(name=_FONT, size=13, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 28


def _write_probe(wb, *, preserve_positions=True):
    # Preserve an existing owned input sheet; refuse an unrelated sheet collision.
    positions = {name: wb.sheetnames.index(name) for name in OWNED if name in wb.sheetnames}
    new_input = "Finance Input" not in wb.sheetnames
    if not new_input:
        inp = wb["Finance Input"]
        if inp["A2"].value != SCHEMA:
            raise ValueError("Finance Input is not an FF-0 probe; refusing replacement")
        for row in range(LAST_ROW + 1, inp.max_row + 1):
            if any(inp.cell(row, c).value is not None for c in range(1, 6)):
                raise ValueError("Input exceeds prepared FF-0 capacity; no rows were discarded")
    else:
        if any(n in wb.sheetnames for n in OWNED[1:]):
            raise ValueError("Existing finance sheet collision")
        inp = wb.create_sheet("Finance Input")
        inp["A1"] = "FINANCE INTEGRATION PROBE"
        inp["A2"] = SCHEMA
        inp["A3"] = "Synthetic values only; this is not a production forecast."
        inp["A4"] = "BOQ basis: Contract Value. No cost-efficiency or Activity P/L."
        settings = [("Gross certificate", 100), ("Certificate date", datetime(2026, 1, 31)),
                    ("Retention rate", .05), ("Credit days", 30), ("Opening cash at 1-Jan", 10),
                    ("Cash actuals through", datetime(2026, 2, 28)), ("Remaining receipt override", None),
                    ("Already received", 40)]
        for row, (label, value) in enumerate(settings, 7):
            inp.cell(row, 1, label); inp.cell(row, 2, value)
        inp["B8"].number_format = inp["B12"].number_format = "dd-mmm-yyyy"
        inp["B9"].number_format = "0.0%"
        inp["A16"] = "Edit blue cells, then F9 / Save. Blank override uses formula; zero cancels."
        inp["A17"] = "20 prepared rows (20-39). No automatic table expansion promised."
        inp["A18"] = "Dates outside Jan-Jun 2026 need review; see Cash Flow warnings."
        for c, h in enumerate(("Event ID", "Cash date", "Kind", "Cash In", "Cash Out"), 1):
            inp.cell(19, c, h)
        events = [("ACT-01", datetime(2026, 1, 15), "Actual", 0, 20),
                  ("ACT-02", datetime(2026, 2, 10), "Actual", 40, 0),
                  ("FCT-01", datetime(2026, 4, 1), "Forecast", 0, 25),
                  ("RET-01", datetime(2026, 6, 30), "Forecast", 5, 0)]
        for row, values in enumerate(events, FIRST_ROW):
            for c, value in enumerate(values, 1):
                inp.cell(row, c, value)
        for row in range(FIRST_ROW, LAST_ROW + 1):
            inp.cell(row, 2).number_format = "dd-mmm-yyyy"
        table = Table(displayName="FF0_CashEvents", ref=f"A19:E{LAST_ROW}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        inp.add_table(table)
        for kind, formula, target in [("list", '"Actual,Forecast"', f"C{FIRST_ROW}:C{LAST_ROW}"),
                                      ("decimal", "0", f"D{FIRST_ROW}:E{LAST_ROW}"),
                                      ("decimal", "0", "B7 B10 B11 B13 B14")]:
            dv = DataValidation(type=kind, operator="greaterThanOrEqual" if kind == "decimal" else None,
                                formula1=formula, allow_blank=True)
            dv.showErrorMessage = True
            inp.add_data_validation(dv); dv.sqref = target
    for n in OWNED[1:]:
        if n in wb.sheetnames:
            if wb[n]["H1"].value != SCHEMA:
                raise ValueError(f"{n} is not an FF-0 probe")
            del wb[n]
    data = wb.create_sheet("Finance_Data")
    cash = wb.create_sheet("Cash Flow")
    for ws in (data, cash):
        ws["H1"] = SCHEMA
    for name, cell in [("FF0_Credit_Days", "B10"), ("FF0_Actual_Through", "B12")]:
        replace_defined_name(wb, name, f"'Finance Input'!${cell[0]}${cell[1:]}")
    data["A1"] = "FORMULA PROBES"
    probes = [("Due date", "='Finance Input'!B8+FF0_Credit_Days"),
              ("Net certificate", "='Finance Input'!B7*(1-'Finance Input'!B9)"),
              ("Retention withheld", "='Finance Input'!B7*'Finance Input'!B9"),
              ("Remaining receipt", '=IF(\'Finance Input\'!B13="",B3-\'Finance Input\'!B14,\'Finance Input\'!B13)')]
    for row, (label, formula) in enumerate(probes, 2):
        data.cell(row, 1, label); data.cell(row, 2, formula)
    data["B2"].number_format = "dd-mmm-yyyy"
    cash["A1"] = "CASH FLOW PROOF (SYNTHETIC)"
    cash["A2"] = "Formula / chart integration only. F9 or Save to calculate."
    cash["A3"] = "Expected default closing cash: 65. See acceptance guide for edits."
    cash["A4"] = "Actual completeness / due date"
    cash["B4"] = "=FF0_Actual_Through"; cash["C4"] = "=Finance_Data!B2"
    cash["B4"].number_format = cash["C4"].number_format = "dd-mmm-yyyy"
    for c, h in enumerate(("Month", "Actual In", "Actual Out", "Forecast In", "Forecast Out", "Closing cash"), 1):
        cash.cell(5, c, h)
    # Fixed small horizon deliberately avoids building an FF-1 financial engine.
    for row, month in enumerate(range(1, 7), 6):
        cash.cell(row, 1, datetime(2026, month, 1)).number_format = "mmm-yyyy"
        start = f"A{row}"
        end = f"DATE(YEAR({start}),MONTH({start})+1,1)"
        dates = "'Finance Input'!$B$20:$B$39"
        kinds = "'Finance Input'!$C$20:$C$39"
        for col, kind, amount_col in [(2, "Actual", "D"), (3, "Actual", "E"),
                                      (4, "Forecast", "D"), (5, "Forecast", "E")]:
            boundary = '"<="' if kind == "Actual" else '">"'
            amount = f"'Finance Input'!${amount_col}$20:${amount_col}$39"
            formula = f'=SUMIFS({amount},{dates},">="&{start},{dates},"<"&{end},{kinds},"{kind}",{dates},{boundary}&FF0_Actual_Through)'
            if col == 4:
                formula += f'+IF(AND(Finance_Data!$B$2>FF0_Actual_Through,Finance_Data!$B$2>={start},Finance_Data!$B$2<{end}),Finance_Data!$B$5,0)'
            cash.cell(row, col, formula)
        previous = "'Finance Input'!B11" if row == 6 else f"F{row-1}"
        cash.cell(row, 6, f"={previous}+B{row}-C{row}+D{row}-E{row}")
    cash["A13"] = "Due outside horizon / overdue"
    cash["B13"] = '=IF(OR(Finance_Data!B2<DATE(2026,1,1),Finance_Data!B2>=DATE(2026,7,1)),"OUT OF RANGE",IF(Finance_Data!B2<=FF0_Actual_Through,"REVIEW OVERDUE","OK"))'
    cash["A14"] = "Events outside prepared horizon"
    cash["B14"] = '=COUNTIFS(\'Finance Input\'!B20:B39,">0",\'Finance Input\'!B20:B39,"<"&DATE(2026,1,1))+COUNTIF(\'Finance Input\'!B20:B39,">="&DATE(2026,7,1))'
    cats = Reference(cash, min_col=1, min_row=6, max_row=11)
    bars = BarChart(); bars.title = "Monthly cash events"; bars.y_axis.title = "Synthetic currency units"
    bars.add_data(Reference(cash, min_col=2, max_col=5, min_row=5, max_row=11), titles_from_data=True)
    bars.set_categories(cats); bars.width = 23; bars.height = 9
    cash.add_chart(bars, "A17")
    line = LineChart(); line.title = "Month-end cash balance"
    line.add_data(Reference(cash, min_col=6, min_row=5, max_row=11), titles_from_data=True)
    line.set_categories(cats); _date_axis_for_line_chart(line, title="Month")
    line.width = 23; line.height = 9
    line.series[0].graphicalProperties.line.solidFill = _COLORS["blue"]
    cash.add_chart(line, "A35")
    for ws in (inp, cash, data):
        if ws is not inp or new_input:
            _style(ws)
        ws.protection.sheet = True
        ws.protection.selectLockedCells = False
        ws.protection.selectUnlockedCells = False
        ws.protection.autoFilter = False
    inp.sheet_state = cash.sheet_state = "visible"; data.sheet_state = "hidden"
    for row in list(inp["B7:B14"]) + list(inp[f"A{FIRST_ROW}:E{LAST_ROW}"]):
        for c in row:
            c.protection = Protection(locked=False)
            c.font = Font(name=_FONT, size=11, color="0000FF")
    if preserve_positions:
        for name, index in sorted(positions.items(), key=lambda item: item[1]):
            wb.move_sheet(name, offset=index - wb.sheetnames.index(name))
    return inp


def extend(source, output, *, restore_opaque=False, preserve_positions=True):
    """Temporary-output extension, preserving existing calculation properties.

    No global finalize call: that would mutate unrelated sheets and hide these
    currently unregistered probe titles. Lifecycle tests explicitly test that gap.
    """
    source, output = Path(source).resolve(), Path(output).resolve()
    if source == output or output.exists():
        raise ValueError("FF-0 requires a new output path")
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("FF-0 probe supports .xlsx only; macro preservation is unproven")
    output.parent.mkdir(parents=True, exist_ok=True)
    before = snapshot(source)
    fd, name = tempfile.mkstemp(suffix=".xlsx", dir=output.parent)
    os.close(fd); temp = Path(name)
    try:
        wb = load_workbook(source, keep_links=True)
        try:
            calc = copy(wb.calculation)
            _write_probe(wb, preserve_positions=preserve_positions)
            reassert_traditional_overlay_transparency(wb)
            wb.calculation = calc
            wb.save(temp)
        finally:
            wb.close()
        if restore_opaque:
            restore_opaque_workbook_parts(source, temp)
        issues = package_issues(temp)
        after = snapshot(temp)
        issues += compare(before, after, [n for n in before["sheets"] if n not in OWNED])
        if before["calculation"] != after["calculation"]:
            issues.append("calculation policy changed")
        for key, value in before["names"].items():
            if not key.startswith("FF0_Credit_") and key != "FF0_Actual_Through" and after["names"].get(key) != value:
                issues.append(f"defined name changed: {key}")
        if issues:
            raise ValueError("Package proof failed: " + "; ".join(issues))
        os.replace(temp, output)
    finally:
        temp.unlink(missing_ok=True)
    return output


def inject_opaque_shape(source, output):
    """A valid DrawingML rectangle sentinel that openpyxl does not model."""
    with ZipFile(source) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sheet = _sheet_parts(parts)["Dashboard"]
    relpath = posixpath.dirname(sheet) + "/_rels/" + posixpath.basename(sheet) + ".rels"
    rel = next(r for r in ET.fromstring(parts[relpath]) if r.get("Type", "").endswith("/drawing"))
    target = rel.get("Target")
    drawing = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(sheet), target))
    root = ET.fromstring(parts[drawing])
    shape = ET.fromstring(f'''<xdr:twoCellAnchor xmlns:xdr="{XDR}" xmlns:a="{A}">
      <xdr:from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>42</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>
      <xdr:to><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>44</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>
      <xdr:sp><xdr:nvSpPr><xdr:cNvPr id="9999" name="FF0_OPAQUE_SENTINEL"/><xdr:cNvSpPr/></xdr:nvSpPr>
      <xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="FFC000"/></a:solidFill></xdr:spPr>
      <xdr:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>FF-0 opaque shape sentinel</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/>
      </xdr:twoCellAnchor>''')
    root.append(shape); parts[drawing] = ET.tostring(root)
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        for n, raw in parts.items():
            z.writestr(n, raw)
    return Path(output)


def has_opaque_shape(path):
    with ZipFile(path) as z:
        return any(b"FF0_OPAQUE_SENTINEL" in z.read(n) for n in z.namelist()
                   if n.startswith("xl/drawings/") and n.endswith(".xml"))


def lifecycle_matrix(source, directory):
    from progress_studio.services.rebuild_service import WorkbookRebuildEngine
    from progress_studio.services.earned_value_rebuild_service import EarnedValueRebuildService
    from progress_studio.services.payment_service import PaymentService
    from progress_studio.services.mapping_store import MappingStore
    from progress_studio.services.workbook_export_service import WorkbookExportService
    from progress_studio.domain.mapping_models import ActivityRow, BOQRow
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    engine = WorkbookRebuildEngine()
    operations = {
        "progress_snapshot": engine.rebuild_progress,
        "progress_live": engine.rebuild_live_progress,
        "payment_snapshot": engine.rebuild_payment,
        "payment_live": engine.rebuild_live_payment,
        "earned_value": EarnedValueRebuildService().generate,
        "payment_breakdown": PaymentService().prepare_payment_breakdown,
    }
    # This remapping probe is only applicable to the known synthetic fixture.
    from progress_studio.infrastructure.excel.rebuild_workbook_reader import RebuildWorkbookReader
    ids = {r.activity_id for r in RebuildWorkbookReader().read_main_dataset(source).activities}
    if ids == {"A1000"}:
        store = MappingStore()
        store.load_activities([ActivityRow("A1000", "1", "1.1", "Concrete")])
        store.load_boq([BOQRow("BOQ|2", "BOQ", 2, "1", "", "", "Concrete", 1000, "BOQ-ONE")])
        store.selected_activity_ids = {"A1000"}; store.selected_boq_ids = {"BOQ|2"}
        store.map_selected(100)
        operations["mapping_export"] = lambda src, out: WorkbookExportService().export(src, out, store)
    before = snapshot(source)
    results = {}
    for name, operation in operations.items():
        output = directory / (name + ".xlsx")
        try:
            with warnings.catch_warnings(record=True) as caught:
                warnings.simplefilter("always", UserWarning)
                operation(source, output)
            after = snapshot(output)
            issues = compare(before, after, OWNED) + package_issues(output)
            for key, value in before["names"].items():
                if key.startswith("FF0_") and after["names"].get(key) != value:
                    issues.append(f"defined name changed: {key}")
            results[name] = {"status": "PASS" if not issues else "FAIL", "issues": issues,
                             "warnings": sorted(set(str(w.message) for w in caught)),
                             "calculation": after["calculation"], "output": output.name}
            refresh = directory / (name + "_finance_refresh.xlsx")
            try:
                extend(output, refresh, restore_opaque=True)
                results[name]["subsequent_finance_refresh"] = "PASS (structural only)"
            except ValueError as exc:
                results[name]["subsequent_finance_refresh"] = "BLOCKED: " + str(exc)
            if name == "earned_value":
                try:
                    extend(output, directory / "ev_append_refresh.xlsx", restore_opaque=True,
                           preserve_positions=False)
                    results[name]["append_order_negative_probe"] = "NO FAILURE OBSERVED"
                except ValueError as exc:
                    results[name]["append_order_negative_probe"] = "BLOCKED: " + str(exc)
        except Exception as exc:
            results[name] = {"status": "ERROR", "issues": [f"{type(exc).__name__}: {exc}"]}
    return results


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n")
